from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
import sys
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def flat(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, list):
        return " | ".join(flat(item) for item in value)
    if isinstance(value, dict):
        return " | ".join(f"{key}: {flat(item)}" for key, item in value.items())
    return str(value)


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    headers: list[str] = []
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)
    if not rows:
        rows, headers = [{"estado": "Sin datos"}], ["estado"]
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: flat(row.get(key)) for key in headers})


def collect(campaign: Path, gold_path: Path, registry: Path, expected_hash: str) -> dict[str, Any]:
    gold = read_json(gold_path)
    cards = {row["group_id"]: row for row in gold.get("cards") or []}
    group_rows: list[dict[str, Any]] = []
    role_rows: list[dict[str, Any]] = []
    field_rows: list[dict[str, Any]] = []
    source_rows: list[dict[str, Any]] = []
    telemetry_rows: list[dict[str, Any]] = []
    error_rows: list[dict[str, Any]] = []
    for phase in ("calibration", "holdout"):
        records_path = campaign / phase / "records.json"
        if not records_path.exists():
            continue
        for record in read_json(records_path):
            group_id = record["group_id"]
            final = record.get("final_decision") or {}
            card = cards.get(group_id, {})
            direction = final.get("direction_assessment") or {}
            gold_checks = {
                "core_status": final.get("core_status") in (card.get("acceptable_core_statuses") or []),
                "inference_ceiling": final.get("inference_ceiling") in (card.get("acceptable_inference_ceilings") or []),
                "dominant_direction": direction.get("dominant_direction") in (card.get("acceptable_directions") or []),
                "context_usable": final.get("context_usable") == card.get("context_usable"),
            }
            role_errors = [
                f"{role}:{error}" for role in ("curator", "critic", "arbiter")
                for error in record.get(f"{role}_errors") or []
            ]
            group_rows.append({
                "fase": phase, "grupo": group_id, "completo": bool(final),
                "acuerdo_curator_critic": record.get("base_pair_agreement"),
                "arbitraje": record.get("adjudicated"),
                "motivos_arbitraje": record.get("arbitration_reasons"),
                "core_status_final": final.get("core_status"),
                "inference_ceiling_final": final.get("inference_ceiling"),
                "direccion_final": direction.get("dominant_direction"),
                "conflicto_material": direction.get("material_conflict"),
                "confianza": final.get("confidence"),
                "provenance_allowlist": record.get("allowlist_provenance"),
                "dentro_gold": all(gold_checks.values()), "errores_roles": role_errors,
            })
            comparisons = {
                "core_status": (card.get("acceptable_core_statuses"), final.get("core_status")),
                "context_usable": (card.get("context_usable"), final.get("context_usable")),
                "inference_ceiling": (card.get("acceptable_inference_ceilings"), final.get("inference_ceiling")),
                "dominant_direction": (card.get("acceptable_directions"), direction.get("dominant_direction")),
                "material_conflict": ("según evidencia directa", direction.get("material_conflict")),
                "limitations": (card.get("required_limitations"), final.get("limitations")),
            }
            for field, (expected, observed) in comparisons.items():
                field_rows.append({
                    "fase": phase, "grupo": group_id, "campo": field,
                    "gold": expected, "sol_final": observed,
                    "coincide": gold_checks.get(field, True),
                })
            for role in ("curator", "critic", "arbiter"):
                semantic = record.get(f"{role}_semantic")
                decision = record.get(f"{role}_decision")
                if not semantic:
                    continue
                role_direction = (decision or {}).get("direction_assessment") or {}
                role_rows.append({
                    "fase": phase, "grupo": group_id, "rol": role,
                    "core_status": (decision or {}).get("core_status"),
                    "inference_ceiling": (decision or {}).get("inference_ceiling"),
                    "direccion": role_direction.get("dominant_direction"),
                    "conflicto_material": role_direction.get("material_conflict"),
                    "confianza": semantic.get("confidence"),
                    "errores": record.get(f"{role}_errors") or [],
                })
                for item in semantic.get("used_evidence") or []:
                    source_rows.append({
                        "fase": phase, "grupo": group_id, "rol": role,
                        "evidence_id": item.get("evidence_id"), "disposición": "usada",
                        "roles": item.get("roles"), "dirección": item.get("supports_direction"),
                        "clase_conflicto": item.get("core_conflict_class"),
                        "rationale": item.get("core_conflict_rationale"),
                    })
                for item in semantic.get("excluded_evidence") or []:
                    source_rows.append({
                        "fase": phase, "grupo": group_id, "rol": role,
                        "evidence_id": item.get("evidence_id"), "disposición": "excluida",
                        "roles": "", "dirección": "", "clase_conflicto": "",
                        "rationale": item.get("reason_code"),
                    })
            for audit in record.get("call_audits") or []:
                for attempt in audit.get("attempts") or []:
                    telemetry_rows.append({
                        "fase": phase, "grupo": group_id, "rol": audit.get("role"),
                        "intento": attempt.get("attempt"), "response_id": attempt.get("response_id"),
                        "modelo_efectivo": attempt.get("effective_model"),
                        "estado": attempt.get("response_status") or attempt.get("http_status"),
                        "input_tokens": attempt.get("input_tokens"),
                        "cached_input_tokens": attempt.get("cached_input_tokens"),
                        "output_tokens": attempt.get("output_tokens"),
                        "reasoning_tokens": attempt.get("reasoning_tokens"),
                        "latencia_segundos": attempt.get("latency_seconds"),
                        "costo_usd": attempt.get("estimated_cost_usd"),
                        "observabilidad_costo": attempt.get("cost_observability"),
                    })
            for role_error in role_errors:
                error_rows.append({"fase": phase, "grupo": group_id, "tipo": "rol", "error": role_error})
    for phase in ("calibration", "holdout"):
        report_path = campaign / phase / "phase_acceptance.json"
        if report_path.exists():
            for error in read_json(report_path).get("errors") or []:
                error_rows.append({"fase": phase, "grupo": error.split(":", 1)[0], "tipo": "gate", "error": error})
    final = read_json(campaign / "final_evaluation_report.json") if (campaign / "final_evaluation_report.json").exists() else {}
    budget = read_json(campaign / "budget_ledger.json") if (campaign / "budget_ledger.json").exists() else {}
    registry_hash = sha256(registry)
    summary = [{
        "campaña": campaign.name, "estado": final.get("status", "incompleto"),
        "resultado": "PASÓ" if final.get("passed") else "FALLÓ",
        "grupos_completos": len(group_rows),
        "acuerdos_base": sum(int(bool(row["acuerdo_curator_critic"])) for row in group_rows),
        "arbitrajes": sum(int(bool(row["arbitraje"])) for row in group_rows),
        "errores": len(error_rows), "costo_total_usd": budget.get("accounted_campaign_cost_usd"),
        "registro_intacto": registry_hash.lower() == expected_hash.lower(),
        "hash_registro": registry_hash,
        "siguiente_estado": "detenido; requiere holdout nuevo no visto" if not final.get("passed") else "apto para preparación 105",
    }]
    return {
        "summary": summary, "groups": group_rows, "roles": role_rows, "fields": field_rows,
        "sources": source_rows, "telemetry": telemetry_rows, "errors": error_rows,
        "final": final, "budget": budget, "cards": cards,
    }


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = styles[name]; style.font.name = "Calibri"; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color); style.font.bold = True
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = "HEAL | Auditoría interna Semantic V2"
    header.runs[0].font.name = "Calibri"; header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor.from_string("667085")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Documento interno - no constituye validación clínica").font.size = Pt(8)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc_pr.append(shd)


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False; table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW")) or OxmlElement("w:tblW")
    if tbl_w.getparent() is None: tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths))); tbl_w.set(qn("w:type"), "dxa")
    ind = tbl_pr.find(qn("w:tblInd")) or OxmlElement("w:tblInd")
    if ind.getparent() is None: tbl_pr.append(ind)
    ind.set(qn("w:w"), "120"); ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcw = cell._tc.get_or_add_tcPr().find(qn("w:tcW")) or OxmlElement("w:tcW")
            if tcw.getparent() is None: cell._tc.get_or_add_tcPr().append(tcw)
            tcw.set(qn("w:w"), str(widths[index])); tcw.set(qn("w:type"), "dxa")


def add_doc_table(doc: Document, headers: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]; cell.text = header; shade(cell, "E8EEF5")
        for run in cell.paragraphs[0].runs: run.bold = True; run.font.size = Pt(8.5)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = flat(value)
            for run in cells[index].paragraphs[0].runs: run.font.size = Pt(8)
    set_table_geometry(table, widths)


def create_docx(path: Path, data: dict[str, Any]) -> None:
    doc = Document(); style_doc(doc)
    title = doc.add_paragraph(); title.paragraph_format.space_after = Pt(4)
    run = title.add_run("AUDITORÍA SEMANTIC V2"); run.bold = True; run.font.size = Pt(23)
    subtitle = doc.add_paragraph("Campaña verification-3 | Calibration aprobada, holdout detenido")
    subtitle.paragraph_format.space_after = Pt(16); subtitle.runs[0].font.size = Pt(13)
    summary = data["summary"][0]
    add_doc_table(doc, ["Campo", "Resultado"], [
        ["Estado formal", summary["estado"]], ["Calibration", "6/6 - PASÓ"],
        ["Holdout", "6/6 ejecutados; FALLÓ 1 gate Gold"],
        ["Acuerdo base", f"{summary['acuerdos_base']}/12"], ["Arbitrajes", summary["arbitrajes"]],
        ["Costo total observado", f"USD {float(summary['costo_total_usd'] or 0):.5f}"],
        ["Registro activo", "Intacto" if summary["registro_intacto"] else "CAMBIÓ"],
    ], [2700, 6660])
    doc.add_heading("Conclusión ejecutiva", level=1)
    doc.add_paragraph(
        "Calibration pasó todos los gates. El holdout completó las seis parejas de curator/critic con acuerdo 6/6, "
        "cero arbitrajes y cero errores de contrato, identidad, cutoff, allowlist o aprobaciones sin respaldo. "
        "El único fallo fue COL14A1:T1.5:gold:inference_ceiling: el normalizador derivó initial_guide_candidate "
        "cuando el Gold firmado permite únicamente context_only."
    )
    doc.add_paragraph(
        "Por el protocolo de holdout único, la campaña quedó sellada como holdout_failed_requires_new_unseen_holdout. "
        "No se prepararon ni ejecutaron los 105 grupos y no debe repetirse este mismo holdout."
    )
    doc.add_heading("Resultado por grupo", level=1)
    add_doc_table(doc, ["Fase", "Grupo", "Acuerdo", "Estado", "Techo", "Gold"], [
        [row["fase"], row["grupo"], "Sí" if row["acuerdo_curator_critic"] else "No",
         row["core_status_final"], row["inference_ceiling_final"], "OK" if row["dentro_gold"] else "FALLA"]
        for row in data["groups"]
    ], [1100, 1700, 900, 1500, 2300, 1860])
    doc.add_heading("Análisis del único fallo", level=1)
    doc.add_paragraph(
        "Curator y critic coincidieron en approved, supports_relation y ausencia de conflicto material. Ambos usaron "
        "evidencia humana y funcional válida, pero esa evidencia sostiene el mecanismo biológico general; no demuestra "
        "aplicabilidad humana directa suficiente para una inferencia individual. El contrato semántico actual no distingue "
        "con precisión esa diferencia y el normalizador elevó automáticamente el techo."
    )
    add_doc_table(doc, ["Elemento", "Gold", "Sol normalizado"], [[
        "Inference ceiling COL14A1", "context_only", "initial_guide_candidate"
    ]], [2600, 3000, 3760])
    doc.add_heading("Qué debe revisarse antes de otra campaña", level=1)
    for text in (
        "Separar evidencia humana contextual de evidencia humana directamente aplicable al mecanismo central.",
        "Evitar que fuentes canónicas humanas activen por sí solas human_applicability para initial_guide_candidate.",
        "Agregar un gate determinístico de aplicabilidad directa (no sólo presencia de roles human + functional).",
        "Definir un nuevo holdout no visto; estos seis grupos ya no pueden reutilizarse como holdout.",
    ):
        doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("Controles operativos", level=1)
    doc.add_paragraph(f"Hash final del registro activo: {summary['hash_registro']}")
    doc.add_paragraph("Luna, VCF, publicación y modificación del registry: no ejecutados.")
    doc.add_paragraph("Los outputs crudos, response IDs y telemetría permanecen en el ZIP técnico separado.")
    doc.save(path)


def create_pdf(path: Path, data: dict[str, Any]) -> None:
    styles = getSampleStyleSheet()
    body = ParagraphStyle("BodyHeal", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5,
                          leading=12.5, spaceAfter=7)
    h1 = ParagraphStyle("H1Heal", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15,
                        textColor=colors.HexColor("#2E74B5"), spaceBefore=13, spaceAfter=7)
    title = ParagraphStyle("TitleHeal", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=22,
                           leading=25, alignment=TA_LEFT, textColor=colors.HexColor("#0B2545"))
    doc = SimpleDocTemplate(str(path), pagesize=letter, rightMargin=0.7*inch, leftMargin=0.7*inch,
                            topMargin=0.7*inch, bottomMargin=0.7*inch)
    story: list[Any] = [Paragraph("AUDITORÍA SEMANTIC V2", title),
                        Paragraph("verification-3 | Resultado formal del holdout único", body), Spacer(1, 8)]
    summary = data["summary"][0]
    summary_table = Table([
        ["Estado", summary["estado"]], ["Calibration", "6/6 - PASÓ"],
        ["Holdout", "6/6 ejecutados; 1 gate Gold falló"], ["Acuerdo", f"{summary['acuerdos_base']}/12"],
        ["Arbitrajes", str(summary["arbitrajes"])], ["Costo", f"USD {float(summary['costo_total_usd'] or 0):.5f}"],
        ["Registro", "Intacto" if summary["registro_intacto"] else "CAMBIÓ"],
    ], colWidths=[1.55*inch, 5.2*inch])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (0,-1), colors.HexColor("#E8EEF5")), ("FONTNAME", (0,0), (0,-1), "Helvetica-Bold"),
        ("FONTNAME", (0,0), (-1,-1), "Helvetica"), ("FONTSIZE", (0,0), (-1,-1), 8.5),
        ("GRID", (0,0), (-1,-1), 0.4, colors.HexColor("#CBD5E1")), ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("LEFTPADDING", (0,0), (-1,-1), 6), ("RIGHTPADDING", (0,0), (-1,-1), 6),
        ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5),
    ])); story.append(summary_table)
    story += [Paragraph("Conclusión ejecutiva", h1), Paragraph(
        "Calibration pasó. Holdout completó 6/6 grupos con acuerdo 6/6, cero arbitrajes y cero errores técnicos. "
        "El único fallo fue COL14A1:T1.5:gold:inference_ceiling: Sol normalizó initial_guide_candidate y el Gold "
        "firmado admite sólo context_only.", body), Paragraph(
        "La campaña quedó sellada como holdout_failed_requires_new_unseen_holdout. No se prepararon ni ejecutaron los 105 grupos.", body)]
    story.append(Paragraph("Resultado por grupo", h1))
    group_table = Table([["Fase", "Grupo", "Acuerdo", "Estado", "Techo", "Gold"]] + [[
        row["fase"], row["grupo"], "Sí" if row["acuerdo_curator_critic"] else "No", row["core_status_final"],
        row["inference_ceiling_final"], "OK" if row["dentro_gold"] else "FALLA"
    ] for row in data["groups"]], repeatRows=1, colWidths=[0.75*inch, 1.15*inch, 0.6*inch, 0.9*inch, 1.55*inch, 0.7*inch])
    group_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#E8EEF5")), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTNAME", (0,1), (-1,-1), "Helvetica"), ("FONTSIZE", (0,0), (-1,-1), 7),
        ("GRID", (0,0), (-1,-1), 0.3, colors.HexColor("#CBD5E1")), ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 4),
        ("BACKGROUND", (0,1), (-1,-1), colors.white),
    ])); story.append(group_table)
    story += [PageBreak(), Paragraph("Causa estructural", h1), Paragraph(
        "El contrato actual distingue human_applicability y functional_compatibility, pero no exige que la evidencia humana "
        "sea directamente aplicable a una inferencia individual. Fuentes canónicas humanas y estudios observacionales "
        "activaron el gate determinístico aunque sólo sustentan contexto biológico.", body),
        Paragraph("Próximo paso permitido", h1), Paragraph(
            "Revisar el gate de aplicabilidad directa y definir un nuevo conjunto holdout no visto. No repetir los seis casos "
            "actuales como holdout y no avanzar a los 105 con este resultado.", body),
        Paragraph("Controles", h1), Paragraph(
            f"Registro activo intacto: {summary['hash_registro']}. Luna, VCF y publicación no ejecutados.", body)]
    doc.build(story)


def zip_tree(zip_path: Path, roots: list[tuple[Path, str]], *, exclude_json: bool) -> None:
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for root, prefix in roots:
            if root.is_file():
                if not (exclude_json and root.suffix.lower() == ".json"):
                    archive.write(root, f"{prefix}/{root.name}" if prefix else root.name)
                continue
            if not root.exists():
                continue
            for path in sorted(root.rglob("*")):
                if path.resolve() == zip_path.resolve():
                    continue
                relative = path.relative_to(root)
                if exclude_json and (
                    path.suffix.lower() in {".json", ".ndjson", ".zip"}
                    or any(part.startswith("_") for part in relative.parts)
                ):
                    continue
                if path.is_file() and not (exclude_json and path.suffix.lower() == ".json"):
                    archive.write(path, f"{prefix}/{relative}".replace("\\", "/"))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--verification-campaign", required=True)
    parser.add_argument("--full-campaign", required=True)
    parser.add_argument("--gold", required=True)
    parser.add_argument("--registry", required=True)
    parser.add_argument("--expected-registry-sha256", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--phase", required=True)
    parser.add_argument("--orchestration-root")
    args = parser.parse_args()
    campaign = Path(args.verification_campaign).resolve(); output = Path(args.output_dir).resolve()
    output.mkdir(parents=True, exist_ok=True)
    data = collect(campaign, Path(args.gold).resolve(), Path(args.registry).resolve(), args.expected_registry_sha256)
    csv_map = {
        "RESUMEN.csv": data["summary"], "RESULTADOS_POR_GRUPO.csv": data["groups"],
        "GOLD_VS_SOL_CAMPO_A_CAMPO.csv": data["fields"], "CURATOR_CRITIC_ARBITER.csv": data["roles"],
        "ERRORES_Y_CONFLICTOS.csv": data["errors"], "FUENTES_USADAS_Y_EXCLUIDAS.csv": data["sources"],
        "TELEMETRIA_Y_COSTOS.csv": data["telemetry"],
    }
    for name, rows in csv_map.items(): write_csv(output / name, rows)
    guide = output / "GUIA_DE_LECTURA.md"
    guide.write_text(
        "# Guía de lectura\n\n1. Empezar por RESUMEN y el informe PDF.\n"
        "2. Revisar RESULTADOS_POR_GRUPO; la única fila fuera del Gold es COL14A1.\n"
        "3. En GOLD_VS_SOL filtrar `coincide = No` para ver el campo exacto.\n"
        "4. CURATOR_CRITIC_ARBITER muestra que ambos roles coincidieron y no hubo árbitro.\n"
        "5. FUENTES permite auditar cada evidence_id permitido, usado o excluido.\n"
        "6. TELEMETRIA contiene response IDs, tokens, latencia y costo; no contiene la API key.\n\n"
        "Resultado: holdout fallido; no repetir estos seis casos como holdout y no ejecutar los 105.\n",
        encoding="utf-8",
    )
    create_docx(output / "Informe_Auditoria_SemanticV2_Verification3.docx", data)
    create_pdf(output / "Informe_Auditoria_SemanticV2_Verification3.pdf", data)
    workbook_data = output / "_workbook_data.json"
    workbook_data.write_text(json.dumps({key: value for key, value in data.items() if key in {"summary","groups","fields","roles","errors","sources","telemetry"}}, ensure_ascii=False), encoding="utf-8")
    manifest_rows = []
    for path in sorted(output.glob("*")):
        if (path.is_file() and path.name not in {"MANIFEST_SHA256.csv", "_workbook_data.json"}
                and path.suffix.lower() not in {".json", ".ndjson", ".zip"}):
            manifest_rows.append({"archivo": path.name, "sha256": sha256(path), "bytes": path.stat().st_size})
    write_csv(output / "MANIFEST_SHA256.csv", manifest_rows)
    human_zip = output / "LLM1_Tier1_SemanticV2_Verification3_HUMANO.zip"
    zip_tree(human_zip, [(output, "")], exclude_json=True)
    technical_zip = output / "LLM1_Tier1_SemanticV2_Verification3_TECNICO.zip"
    technical_roots = [(campaign, "campaign"), (Path(args.full_campaign).resolve(), "full_campaign")]
    if args.orchestration_root:
        technical_roots.append((Path(args.orchestration_root).resolve(), "orchestration"))
    zip_tree(technical_zip, technical_roots, exclude_json=False)
    print(json.dumps({"status": "audit_package_created", "output": str(output), "phase": args.phase}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
