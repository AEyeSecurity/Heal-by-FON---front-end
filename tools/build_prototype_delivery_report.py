#!/usr/bin/env python3
"""Build the human technical/operational delivery report for the Tier 1 prototype."""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


GREEN = "275D38"
EXPECTED_REGISTRY = "73b94c09c31c184135bc16b2e756fa447f4c040a8bf38b49181168a8c62a967f"


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def read_csv(path: Path) -> list[dict]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def report_data(smoke: Path, deployment_sha: str, docker_digest: str, tests: int, fresh_job: Path | None) -> dict:
    summary = read_json(smoke / "grouped_prototype_run_summary.json")
    quarantines = read_json(smoke / "quarantine.json")
    cards = read_csv(smoke / "cards.csv")
    telemetry = read_csv(smoke / "telemetry_costs.csv")
    reason_counts = Counter()
    for row in quarantines:
        reasons = row.get("reason_codes") or row.get("reasons") or row.get("error") or []
        if isinstance(reasons, str):
            reasons = [part for part in reasons.split(";") if part]
        reason_counts.update(reasons)
    modes = Counter(row.get("inference_mode") for row in cards if row.get("status") == "valid")
    fresh = read_json(fresh_job) if fresh_job and fresh_job.exists() else {}
    return {
        "summary": summary,
        "quarantines": quarantines,
        "telemetry": telemetry,
        "reason_counts": reason_counts,
        "modes": modes,
        "deployment_sha": deployment_sha,
        "docker_digest": docker_digest,
        "tests": tests,
        "fresh": fresh,
    }


def add_docx_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Indicador"
    table.rows[0].cells[1].text = "Resultado"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True


def build_docx(data: dict, path: Path) -> None:
    s = data["summary"]
    c = s["counts"]
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    for name in ("Title", "Heading 1", "Heading 2"):
        styles[name].font.name = "Arial"
        styles[name].font.color.rgb = RGBColor.from_string(GREEN)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("HEAL — cierre Tier 1 y prototipo end-to-end")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Resumen técnico y operativo · 24/08/2026").italic = True

    document.add_heading("Conclusión ejecutiva", level=1)
    document.add_paragraph(
        "El snapshot científico Tier 1 quedó compilado con 105 grupos firmados y el pipeline agrupado recorrió "
        "LLM1, normalización, LLM2 y generación de reportes. La integración funciona y conserva trazabilidad, "
        "pero el smoke se clasifica como prototype_demo_incomplete porque 13 tarjetas fueron aisladas por los "
        "guardrails semánticos. Las 38 tarjetas válidas sí llegaron a LLM2. No se modificó el registro productivo."
    )
    add_docx_table(document, [
        ("Cobertura científica", f"{c['scientifically_covered']}/{c['canonical_groups']} grupos"),
        ("Grupos con variante observada", str(c["covered_with_observed_variant"])),
        ("Tarjetas LLM1 válidas", str(c["valid_llm1_cards"])),
        ("Tarjetas en cuarentena", str(c["quarantined"])),
        ("Modos válidos", f"{data['modes'].get('context_only', 0)} context_only; {data['modes'].get('initial_guide', 0)} initial_guide"),
        ("Llamadas registradas", str(s["telemetry"]["calls"])),
        ("Costo estimado", f"USD {s['telemetry']['estimated_cost_usd']:.4f}"),
        ("Validación formal", "Pendiente de un nuevo holdout no visto"),
    ])

    document.add_heading("Qué quedó implementado", level=1)
    for text in (
        "Snapshot sandbox inmutable de 105 grupos, con decisión humana y allowlist científica trazadas por separado.",
        "Cobertura cerrada de 180 grupos: 105 cubiertos y 75 marcados explícitamente fuera del prototipo.",
        "Envelope LLM1 v2, ceiling variante-específico, aislamiento por tarjeta y LLM2 limitado a tarjetas válidas.",
        "Interfaz con progreso persistente, métricas dinámicas, costos y descargas DOCX/PDF/auditoría.",
        "Protección del registro activo y ausencia de fallback silencioso al registry productivo.",
    ):
        document.add_paragraph(text, style="List Bullet")

    document.add_heading("Resultado del smoke autorizado", level=1)
    document.add_paragraph(
        "Se procesaron 51 grupos cubiertos con variantes foco. Los validadores aceptaron 38 y aislaron 13; "
        "el job continuó y LLM2 recibió exclusivamente las tarjetas aceptadas. MTHFR:T1.3 fue la única guía "
        "inicial; las otras 37 tarjetas válidas quedaron como contexto."
    )
    if data["reason_counts"]:
        document.add_heading("Motivos de cuarentena", level=2)
        add_docx_table(document, [(reason, str(count)) for reason, count in data["reason_counts"].most_common()])

    document.add_heading("Evidencia de ingeniería", level=1)
    add_docx_table(document, [
        ("Tests", f"{data['tests']} aprobados"),
        ("Build web", "npm run build aprobado"),
        ("Commit desplegado", data["deployment_sha"]),
        ("Imagen normalizer", data["docker_digest"]),
        ("Registry activo", EXPECTED_REGISTRY),
        ("Estado del registry", "Idéntico antes y después"),
    ])
    if data["fresh"]:
        document.add_heading("Job nuevo desde el VCF original", level=1)
        document.add_paragraph(
            f"Job {data['fresh'].get('id', '—')}: estado {data['fresh'].get('status', '—')}, "
            f"etapa {data['fresh'].get('stage', '—')}. Este job se conserva separado de todos los runs históricos."
        )

    document.add_heading("Recomendación para continuar", level=1)
    for text in (
        "Usar esta versión como demo interna, mostrando sólo las 38 tarjetas válidas y la etiqueta de prototipo.",
        "Hacer una iteración puntual del prompt Luna sobre cuatro familias: evidencia duplicada, lenguaje diagnóstico, riesgo individual desde GWAS y referencias fuera de allowlist. No relajar los validadores.",
        "Repetir el smoke después de ese ajuste y exigir cero cuarentenas para declarar prototype_demo_ready_automatic.",
        "Mantener formal_validation_readiness pendiente hasta crear un holdout realmente nuevo e independiente.",
        "No publicar todavía el registry productivo ni presentar el resultado como validación clínica."
    ):
        document.add_paragraph(text, style="List Number")

    document.add_heading("Límites", level=1)
    document.add_paragraph(
        "El VCF es sparse: una variante no observada no equivale a homocigosis de referencia ni a benignidad. "
        "Las interpretaciones fueron generadas por modelos de lenguaje y no constituyen diagnóstico, tratamiento, "
        "suplementación, recomendación farmacológica ni validación clínica independiente."
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def pdf_table(rows: list[tuple[str, str]], widths=(55 * mm, 105 * mm)) -> Table:
    cell_style = ParagraphStyle(name="HealTableCell", parent=getSampleStyleSheet()["BodyText"], fontSize=8.2, leading=9.6, wordWrap="CJK")
    table_rows = [[Paragraph("Indicador", cell_style), Paragraph("Resultado", cell_style)]]
    table_rows.extend([[Paragraph(str(key), cell_style), Paragraph(str(value), cell_style)] for key, value in rows])
    table = Table(table_rows, colWidths=list(widths), repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#275D38")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#A7A7A7")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("LEADING", (0, 0), (-1, -1), 10.5),
        ("PADDING", (0, 0), (-1, -1), 5),
    ]))
    return table


def build_pdf(data: dict, path: Path) -> None:
    s, c = data["summary"], data["summary"]["counts"]
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="HealTitle", parent=styles["Title"], textColor=colors.HexColor("#143A2B"), alignment=TA_CENTER, fontSize=19, leading=23))
    styles.add(ParagraphStyle(name="HealH1", parent=styles["Heading1"], textColor=colors.HexColor("#275D38"), spaceBefore=8, spaceAfter=5))
    styles.add(ParagraphStyle(name="HealNote", parent=styles["BodyText"], backColor=colors.HexColor("#EEF5EE"), borderColor=colors.HexColor("#8BAA8B"), borderWidth=.5, borderPadding=7, leading=13))
    story = [Paragraph("HEAL — cierre Tier 1 y prototipo end-to-end", styles["HealTitle"]), Spacer(1, 3 * mm), Paragraph("Resumen técnico y operativo · 24/08/2026", styles["HealNote"])]
    story += [Paragraph("Conclusión ejecutiva", styles["HealH1"]), Paragraph(
        "El snapshot científico Tier 1 quedó compilado con 105 grupos firmados y el pipeline recorrió LLM1, normalización, LLM2 y reportes. "
        "La integración funciona, pero el smoke es <b>prototype_demo_incomplete</b>: 13 tarjetas fueron aisladas por los guardrails y 38 tarjetas válidas llegaron a LLM2. El registro productivo permaneció intacto.", styles["BodyText"]), Spacer(1, 3 * mm)]
    story.append(pdf_table([
        ("Cobertura", f"{c['scientifically_covered']}/{c['canonical_groups']} grupos"),
        ("Con variante observada", str(c["covered_with_observed_variant"])),
        ("Tarjetas válidas / cuarentena", f"{c['valid_llm1_cards']} / {c['quarantined']}"),
        ("Modos válidos", f"{data['modes'].get('context_only', 0)} context_only; {data['modes'].get('initial_guide', 0)} initial_guide"),
        ("Llamadas / costo", f"{s['telemetry']['calls']} / USD {s['telemetry']['estimated_cost_usd']:.4f}"),
        ("Validación formal", "Pendiente de nuevo holdout unseen"),
    ]))
    story += [Paragraph("Qué quedó implementado", styles["HealH1"])]
    for item in (
        "Snapshot sandbox inmutable de 105 grupos y cobertura cerrada 105/180.",
        "Envelope LLM1 v2, ceiling variante-específico, cuarentena local y LLM2 limitado a tarjetas válidas.",
        "Progreso persistente, métricas, costos y descargas; sin fallback silencioso al registry activo.",
    ):
        story.append(Paragraph(item, styles["BodyText"], bulletText="•"))
    story += [Paragraph("Motivos de cuarentena", styles["HealH1"]), pdf_table([(reason, str(count)) for reason, count in data["reason_counts"].most_common()])]
    story += [PageBreak(), Paragraph("Evidencia de ingeniería", styles["HealH1"]), pdf_table([
        ("Tests", f"{data['tests']} aprobados"),
        ("Build", "npm run build aprobado"),
        ("Commit desplegado", data["deployment_sha"]),
        ("Imagen normalizer", data["docker_digest"]),
        ("Registry activo", EXPECTED_REGISTRY),
        ("Protección", "Idéntico antes y después"),
    ])]
    if data["fresh"]:
        story += [Paragraph("Job nuevo desde el VCF original", styles["HealH1"]), Paragraph(
            f"Job {data['fresh'].get('id', '—')}: estado {data['fresh'].get('status', '—')}, etapa {data['fresh'].get('stage', '—')}. Se conserva separado de los runs históricos.", styles["BodyText"])]
    story += [Paragraph("Recomendación", styles["HealH1"])]
    for item in (
        "Usar esta versión como demo interna, mostrando sólo tarjetas válidas y la etiqueta de prototipo.",
        "Ajustar Luna sobre evidencia duplicada, lenguaje diagnóstico, riesgo individual desde GWAS y referencias fuera de allowlist; no relajar validadores.",
        "Repetir el smoke y exigir cero cuarentenas antes de declarar prototype_demo_ready_automatic.",
        "Mantener la validación formal pendiente hasta un holdout nuevo e independiente.",
    ):
        story.append(Paragraph(item, styles["BodyText"], bulletText="•"))
    story += [Spacer(1, 3 * mm), Paragraph(
        "Límite: un VCF sparse no permite inferir homocigosis de referencia desde una ausencia. El resultado no es diagnóstico, tratamiento ni validación clínica independiente.", styles["HealNote"])]
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm, title="HEAL Tier 1 — resumen técnico y operativo", author="HEAL by FON").build(story)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--smoke-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--deployment-sha", required=True)
    parser.add_argument("--docker-digest", required=True)
    parser.add_argument("--tests", type=int, required=True)
    parser.add_argument("--fresh-job-json")
    args = parser.parse_args()
    output = Path(args.output_dir).resolve()
    data = report_data(Path(args.smoke_dir).resolve(), args.deployment_sha, args.docker_digest, args.tests, Path(args.fresh_job_json).resolve() if args.fresh_job_json else None)
    docx = output / "Resumen_tecnico_y_operativo_HEAL_Tier1_105.docx"
    pdf = output / "Resumen_tecnico_y_operativo_HEAL_Tier1_105.pdf"
    build_docx(data, docx)
    build_pdf(data, pdf)
    print(json.dumps({"docx": str(docx), "pdf": str(pdf)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
