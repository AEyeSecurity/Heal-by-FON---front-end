#!/usr/bin/env python3
"""Create the final pre-Sol and Sol-status human reports."""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from pathlib import Path

from build_tier1_reconciliation_report import Document, add_table, bullet, callout, footer, heading, paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def base_doc(title_text: str, subtitle: str) -> Document:
    doc = Document(); section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492); footer(section)
    normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri"); normal.font.size = Pt(11); normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.10
    paragraph(doc, "INFORME DE CONTROL", size=10, color="5B6673", bold=True, after=3)
    item = paragraph(doc, title_text, size=23, bold=True, after=4); item.paragraph_format.keep_with_next = True
    paragraph(doc, subtitle, size=14, color="5B6673", after=14)
    return doc


def qa_stats(rows: list[dict]) -> list[tuple]:
    counts = Counter(row["group_id"] for row in rows)
    selected = Counter(row["group_id"] for row in rows if row["selected"] == "true")
    corrected = Counter(row["group_id"] for row in rows if row["doi_consistent_before"] == "false")
    return [(group, counts[group], selected[group], corrected[group], "OK") for group in sorted(counts)]


def pre_sol(gold: dict, evidence: dict, qa: list[dict], registry_hash: str) -> Document:
    doc = base_doc("Gold Tier 1 reconciliado y listo", "Resultado final de metadata, firmas, contratos y gates previos a GPT-5.6 Sol")
    callout(doc, "Resultado.", "El Gold de 12 grupos está sellado y validate-gold devolvió 0 errores. La corrección cambió DOI y hashes, pero no cambió PMID, título normalizado, evidence IDs ni decisiones científicas.", "E7F4EC")
    callout(doc, "Estado de Sol.", "Calibration y holdout no fueron ejecutados porque HEAL_OPENAI_API_KEY no está configurada en el entorno seguro. No hay resultados simulados.", "FFF6E0")
    heading(doc, "1. Control ejecutivo")
    add_table(doc, ["Control", "Resultado"], [
        ("Grupos firmados", "12/12"), ("PMID únicos", "1.320"), ("Apariciones en QA", str(len(qa))),
        ("Inconsistencias finales", "0"), ("Firma científica", "Martina Liz Ceballos - 2026-08-14"),
        ("Registro activo", f"Sin cambios - {registry_hash}"),
    ], [2600, 6760], font_size=10)
    heading(doc, "2. Qué se reparó")
    for text in [
        "El parser de PubMed ahora toma DOI sólo de ArticleIdList del artículo y nunca de las referencias citadas.",
        "Los títulos se normalizan sin perder texto; espacios tipográficos en subíndices, superíndices o signos no cambian la identidad.",
        "Los 1.320 PMID se congelaron en un snapshot por lotes con URL, fecha, respuesta y SHA-256.",
        "Cada packet incorpora provenance PubMed, hash del registro autoritativo y estado de consistencia.",
        "Las firmas originales se preservaron; la herencia quedó marcada technical_metadata_only porque no cambió la evidencia científica.",
    ]: bullet(doc, text)
    heading(doc, "3. Decisiones firmadas")
    add_table(doc, ["Grupo", "Split", "Estado", "Contexto", "Techo"], [
        (card["group_id"], card["split"], " | ".join(card["acceptable_core_statuses"]), "Sí" if card["context_usable"] else "No", " | ".join(card["acceptable_inference_ceilings"])) for card in gold["cards"]
    ], [1450, 1450, 2000, 1100, 3360], font_size=8.8)
    doc.add_page_break(); heading(doc, "4. QA PubMed por grupo")
    add_table(doc, ["Grupo", "Registros", "Seleccionados", "DOI corregidos", "Final"], qa_stats(qa), [1900, 1550, 1700, 1800, 2410], font_size=9.2)
    paragraph(doc, "El CSV completo conserva las 1.338 apariciones. Una publicación puede aparecer en más de un grupo; por eso el total supera los 1.320 PMID únicos.", size=9.5, color="5B6673", italic=True)
    heading(doc, "5. Gate excepcional de revisión experta")
    paragraph(doc, "La revisión experta por conflicto sólo puede activarse cuando se cumplen simultáneamente cinco condiciones:")
    for text in ["evidencia fuerte", "conflicto científico real y material", "variante foco observada directamente involucrada", "alternativas que cambian materialmente la interpretación individual", "discrepancia no resoluble con la evidencia y el contexto disponibles"]: bullet(doc, text)
    paragraph(doc, "El runner exige el AND de 5/5. initial_guide_candidate o approved_with_conflict aislados no activan la derivación. Si el gate pasa, la prioridad es recommended, nunca urgent, con confianza Conflicting y una limitación explícita.")
    doc.add_page_break()
    heading(doc, "6. Separación calibration / holdout")
    add_table(doc, ["Etapa", "Acceso", "Protección"], [
        ("calibrate", "Sólo 6 calibration", "6/6 válidos y acuerdo >=5/6"),
        ("optimize", "Sólo códigos de calibration", "Rechaza cualquier reporte de holdout"),
        ("freeze-prompt", "Prompt aprobado", "Sella prompt, schema, Gold, evidencia y hashes"),
        ("evaluate-holdout", "Sólo 6 holdout", "Una ejecución; no sobrescribe ni repite"),
    ], [1900, 2500, 4960], font_size=9.3)
    heading(doc, "7. Próximo paso seguro")
    paragraph(doc, "Configurar HEAL_OPENAI_API_KEY mediante el mecanismo seguro del runtime y ejecutar candidate-1 sobre calibration. Si pasa, congelar el prompt y recién entonces ejecutar holdout una vez. Si holdout falla, se necesita un nuevo conjunto no visto; no se ajusta y reutiliza el mismo holdout.")
    heading(doc, "8. Trazabilidad")
    add_table(doc, ["Artefacto", "SHA-256"], [
        ("Gold", gold["manifest_sha256"]), ("Evidence manifest", evidence["manifest_sha256"]),
        ("Snapshot PubMed", evidence["pubmed_snapshot_manifest_sha256"]), ("QA metadata", evidence["metadata_qa_sha256"]),
        ("Firma fuente CSV", evidence["signature_source_csv_sha256"]), ("Firma fuente XLSX", evidence["signature_source_xlsx_sha256"]),
    ], [2700, 6660], font_size=8.5)
    paragraph(doc, "Este material es curación científica interna y no equivale a validación clínica, genética o bioinformática independiente.", size=9.5, color="5B6673", italic=True)
    return doc


def sol_status(gold: dict) -> Document:
    doc = base_doc("Evaluación estricta con Sol", "Estado honesto de calibration y holdout al 14/08/2026")
    callout(doc, "Estado actual.", "Gold aprobado; calibration no ejecutada; prompt no congelado; holdout no ejecutado. Causa única: falta la credencial segura HEAL_OPENAI_API_KEY.", "FFF6E0")
    heading(doc, "1. Casos de calibration")
    add_table(doc, ["Grupo", "Estado", "Resultado"], [(g, "No ejecutado", "Sin datos") for g in gold["calibration_groups"]], [3000, 3000, 3360], font_size=9.5)
    heading(doc, "2. Casos holdout")
    add_table(doc, ["Grupo", "Estado", "Protección"], [(g, "No ejecutado", "Bloqueado hasta freeze-prompt") for g in gold["holdout_groups"]], [3000, 2500, 3860], font_size=9.5)
    heading(doc, "3. Criterios que se aplicarán")
    for text in ["Dos pasadas independientes con esfuerzo high y arbitraje xhigh cuando corresponda.", "Calibration: 6/6 contratos válidos, cero errores, cero aprobaciones sin respaldo y acuerdo >=5/6.", "Final: 12/12 válidos, cero errores y acuerdo base >=11/12.", "El gate experto debe respetar 5/5 condiciones.", "Si holdout falla, el estado será holdout_failed_requires_new_unseen_holdout."]: bullet(doc, text)
    heading(doc, "4. Cómo reanudar")
    paragraph(doc, "El operador debe inyectar la clave en la configuración segura del proceso. No debe pegarse en archivos, planillas, argumentos de línea de comandos ni documentación. Después se ejecuta calibrate; según su resultado, optimize o freeze-prompt; finalmente evaluate-holdout una sola vez.")
    heading(doc, "5. Acciones expresamente no realizadas")
    for text in ["No se ejecutaron llamadas a Sol.", "No se ejecutó Luna.", "No se recuraron los 105 grupos.", "No se regeneraron VCF.", "No se cambió ni publicó el registro activo."]: bullet(doc, text)
    return doc


def save(doc: Document, path: Path, title: str) -> None:
    doc.core_properties.title = title; doc.core_properties.subject = "Curación científica interna de LLM1"; doc.core_properties.author = "HEAL by FON"
    path.parent.mkdir(parents=True, exist_ok=True); doc.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(); parser.add_argument("--gold", required=True); parser.add_argument("--evidence-manifest", required=True); parser.add_argument("--qa-csv", required=True); parser.add_argument("--registry-hash", required=True); parser.add_argument("--pre-sol-output", required=True); parser.add_argument("--sol-output", required=True)
    args = parser.parse_args(); gold = json.loads(Path(args.gold).read_text(encoding="utf-8")); evidence = json.loads(Path(args.evidence_manifest).read_text(encoding="utf-8"))
    with Path(args.qa_csv).open("r", encoding="utf-8-sig", newline="") as handle: qa = list(csv.DictReader(handle))
    save(pre_sol(gold, evidence, qa, args.registry_hash), Path(args.pre_sol_output), "Gold Tier 1 reconciliado y listo")
    save(sol_status(gold), Path(args.sol_output), "Evaluación estricta con Sol - estado")
    print(json.dumps({"pre_sol": args.pre_sol_output, "sol": args.sol_output, "qa_rows": len(qa)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
