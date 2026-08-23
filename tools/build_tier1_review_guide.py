#!/usr/bin/env python3
"""Create a reader-friendly guide for the Tier 1 gold-standard review."""

from __future__ import annotations

import argparse
import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5B6673"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CAUTION = "FFF6E0"
WHITE = "FFFFFF"
PAGE_WIDTH_DXA = 9360


def set_run_font(run, size=11, color="000000", bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margin(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_dxa, indent=120):
    table.autofit = False
    table_pr = table._tbl.tblPr
    tbl_w = table_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = table_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        table_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for index, width in enumerate(widths_dxa):
        grid.gridCol_lst[index].set(qn("w:w"), str(width))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.tcPr.tcW
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margin(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_para(doc, text="", *, style=None, before=0, after=6, line=1.10, color="000000", size=11, bold=False, italic=False, align=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.17
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        set_run_font(paragraph.add_run(item), size=11)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.17
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        set_run_font(paragraph.add_run(item), size=11)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.space_before = Pt({1: 16, 2: 12, 3: 8}[level])
    paragraph.paragraph_format.space_after = Pt({1: 8, 2: 6, 3: 4}[level])
    run = paragraph.add_run(text)
    set_run_font(run, size={1: 16, 2: 13, 3: 12}[level], color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level], bold=True)
    return paragraph


def add_callout(doc, label, body, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [PAGE_WIDTH_DXA])
    # A callout is one thought. Keep it intact instead of spilling its final
    # lines onto the next page.
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    set_run_font(paragraph.add_run(label + " "), size=11, color=NAVY, bold=True)
    set_run_font(paragraph.add_run(body), size=11, color="000000")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        set_run_font(cell.paragraphs[0].add_run(header), size=9.5, color=NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].paragraphs[0].paragraph_format.space_after = Pt(0)
            set_run_font(cells[index].paragraphs[0].add_run(str(value)), size=9.3)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def section_footer(section):
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    set_run_font(paragraph.add_run("HEAL by FON | Revisión interna Tier 1"), size=8.5, color=MUTED)


def build_document(gold_manifest: dict, review_cards: dict) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    add_para(doc, "GUÍA DE REVISIÓN", size=10, color=MUTED, bold=True, after=3)
    title = add_para(doc, "Recuración Tier 1: qué leer y cómo validar", size=24, color=NAVY, bold=True, after=5)
    title.paragraph_format.keep_with_next = True
    add_para(doc, "Material práctico para revisión interna y conversación con un genetista", size=13, color=MUTED, after=12)
    add_para(doc, f"Fecha: {date.today().strftime('%d/%m/%Y')} | Estado: Gold pendiente de aprobación", size=10, color=MUTED, after=16)

    add_callout(doc, "Objetivo.", "Validar si la relación entre cada gen y su módulo tiene evidencia suficiente como contexto biológico. No se revisan diagnósticos, tratamientos ni resultados individuales de pacientes.")
    add_callout(doc, "Límite importante.", "La aprobación del gold no activa LLM1 ni modifica el registro activo. Sólo define qué comportamiento sería aceptable para calibrar el proceso.", fill=CAUTION)

    add_heading(doc, "1. Resumen en un minuto")
    add_bullets(doc, [
        "Hay 12 fichas de revisión: seis se usarán para calibrar el proceso y seis se guardarán para comprobar que el proceso no se adaptó a los ejemplos.",
        f"La evidencia publicada debe ser del {gold_manifest['evidence_cutoff']} o anterior. Una fuente posterior se conserva como registro, pero no se puede usar para decidir.",
        "Cada ficha contiene hasta 20 fuentes seleccionadas para lectura eficiente. El resto queda registrado y puede revisarse si aparece una duda.",
        "El estado actual es pendiente: todavía no se ejecutó GPT-5.6 Sol para curar ni Luna para interpretar grupos.",
    ])

    add_heading(doc, "2. Archivos que conviene leer", level=1)
    files_rows = [
        ("1. Esta guía", "Lectura inicial", "Da el orden, definiciones y checklist. Usala para no perderte en los JSON."),
        ("2. gold_review_cards.json", "Fichas de evidencia", "Abrir primero. Mirar gen, módulo, propósito, exclusiones y las fuentes seleccionadas."),
        ("3. gold_manifest.json", "Planilla de decisión", "Completar después de mirar la ficha: estado aceptable, contexto, techo de inferencia, fuentes válidas/no válidas y límites."),
        ("4. evidence_manifest.json", "Control técnico", "Corroborar que hay 12/12 paquetes, cutoff correcto, hashes y que no hubo fallos que oculten información."),
        ("5. packets/<GEN>__<TIER>.json", "Detalle para dudas", "Abrir sólo si necesitás discutir una fuente, exclusión, alias del gen o por qué una publicación no fue seleccionada."),
        ("6. Reporte de implementación", "Contexto del sistema", "Explica por qué se reemplazó el 99% withheld y qué salvaguardas quedaron incorporadas."),
    ]
    add_table(doc, ["Archivo", "Para qué sirve", "Qué buscar"], files_rows, [2250, 2100, 5010])

    add_heading(doc, "3. Orden práctico de análisis")
    add_numbered(doc, [
        "Elegí una ficha en gold_review_cards.json. Confirmá que el gen y el módulo tengan sentido juntos. Leé el propósito y, sobre todo, las exclusiones.",
        "Revisá las fuentes seleccionadas. Pregunta clave: ¿esta fuente demuestra una relación con el módulo o sólo menciona el gen en otro contexto?",
        "Marcá fuentes débiles, indirectas, fuera de población/etapa relevante, de método incompleto o claramente ajenas al módulo como no válidas para la decisión.",
        "Definí el resultado conservador en gold_manifest.json. Si hay poca evidencia, el resultado correcto es withheld; no significa que el gen sea benigno ni irrelevante.",
        "Anotá límites claros. Si hay una discrepancia real entre estudios, no fuerces una dirección única.",
        "Repetí para las 12 fichas. Recién cuando todas tengan revisor, fecha y aprobación se puede validar el gold.",
    ])

    add_heading(doc, "4. Cómo elegir el estado", level=1)
    status_rows = [
        ("withheld", "No alcanza la evidencia para sostener una relación útil.", "No equivale a normal, benigno o sin importancia."),
        ("rejected", "Hay evidencia positiva de que el mapeo es incorrecto, el gen no corresponde al módulo o existe una contradicción material.", "No usar sólo porque faltan papers."),
        ("approved", "Existe relación biológica directa y evidencia mínima suficiente.", "Permite contexto; no implica conclusión individual."),
        ("approved_with_conflict", "Hay evidencia suficiente, pero estudios relevantes discrepan.", "La discrepancia debe quedar explícita."),
    ]
    add_table(doc, ["Estado", "Cuándo elegirlo", "Cuidado"], status_rows, [1750, 4750, 2860])

    add_heading(doc, "5. Dos decisiones distintas que no hay que mezclar", level=1)
    add_table(doc, ["Pregunta", "Opciones", "Regla sencilla"], [
        ("¿Hay contexto biológico utilizable?", "Sí / No", "Sólo puede ser Sí con un estado aprobado o aprobado con conflicto."),
        ("¿Hasta dónde puede llegar una futura interpretación?", "none / context_only / initial_guide_candidate", "initial_guide_candidate exige evidencia humana aplicable y evidencia funcional compatible. Aun así, después se validará el genotipo observado."),
    ], [2800, 2450, 4110])

    # Keep the complete examples table together. Splitting its first row across
    # pages makes the practical examples harder to scan during a review meeting.
    doc.add_page_break()
    add_heading(doc, "6. Ejemplos de lectura")
    card_by_id = {card["group_id"]: card for card in review_cards["cards"]}
    examples = []
    for group_id, lesson in [
        ("MTHFR:T1.1", "Una fuente puede mencionar el gen de forma convincente pero en un contexto ajeno al objetivo. Pregunta si el diseño respalda el módulo específico, no sólo si el gen aparece en el título."),
        ("FADS1:T1.3", "La ficha registra un fallo técnico de UniProt. Debe anotarse como limitación; no es evidencia contra el gen ni justifica rechazarlo."),
        ("CYCS:T1.1", "Hay una publicación posterior al cutoff preservada como excluida. No se la debe usar, aunque parezca interesante."),
    ]:
        card = card_by_id[group_id]
        source = next((item for item in card["selected_sources"] if item["source_kind"] == "primary_publication"), None)
        examples.append((group_id, card["module"]["name"], source["evidence_id"] if source else "-", lesson))
    add_table(doc, ["Ficha", "Módulo", "Ejemplo de fuente", "Qué validar"], examples, [1300, 2250, 1500, 4310])
    add_para(doc, "Los ejemplos no son conclusiones científicas sobre esos genes. Sirven para mostrar el tipo de pregunta que debe hacerse al revisar una fuente.", size=9.5, color=MUTED, italic=True, after=8)

    add_heading(doc, "7. Checklist por ficha")
    add_bullets(doc, [
        "Identidad: ¿el símbolo del gen y sus aliases son correctos?",
        "Encaje: ¿la relación gen-módulo es directa o sólo una asociación amplia?",
        "Evidencia: ¿hay al menos una fuente primaria pertinente? ¿la base canónica se usa como contexto y no como única prueba?",
        "Independencia: ¿son estudios realmente distintos o publicaciones derivadas de la misma cohorte?",
        "Dirección: si hay hallazgos opuestos, ¿la diferencia parece real y está suficientemente explicada?",
        "Límites: ¿quedó escrito qué no puede afirmarse?",
        "Alcance: ¿corresponde sólo contexto o hay suficiente base para un candidato a guía inicial?",
        "Trazabilidad: ¿quedan registradas las fuentes válidas y las descartadas?",
    ])

    add_heading(doc, "8. Cuándo pedir una segunda mirada genética")
    add_para(doc, "No hace falta un genetista para cada fila. Conviene involucrarlo cuando haya una de estas dudas:")
    add_bullets(doc, [
        "No está clara la identidad del gen o existe un alias histórico que cambia el mapeo.",
        "El estudio habla de una enfermedad, tumor o modelo experimental y no es obvio que pueda sostener el módulo general.",
        "Dos estudios importantes se contradicen o parecen venir de la misma cohorte.",
        "Se quiere habilitar initial_guide_candidate.",
        "El revisor considera rejected, porque ese estado necesita evidencia positiva, no sólo incertidumbre.",
    ])
    add_callout(doc, "Preguntas útiles para el genetista.", "¿El gen está bien identificado? ¿La relación con este módulo es directa? ¿Qué fuentes son realmente pertinentes? ¿Qué no debería inferirse? ¿Hay evidencia humana y funcional compatible para permitir una guía inicial?", fill=LIGHT_GRAY)

    add_heading(doc, "9. Grupos incluidos en el gold")
    labels = {
        "T1.1": "Sistemas fundamentales y metabolismo de un carbono",
        "T1.2": "Sueño y biología circadiana",
        "T1.3": "Nutrientes y cofactores",
        "T1.4": "Inflamación y tono inmune",
        "T1.5": "Tejido conectivo y resiliencia física",
        "T1.6": "Xenobióticos y estrés oxidativo",
    }
    group_rows = []
    for card in review_cards["cards"]:
        group_rows.append((card["group_id"], labels[card["module"]["module_id"]], "Calibración" if card["split"] == "calibration" else "Control independiente"))
    add_table(doc, ["Grupo", "Eje", "Uso"], group_rows, [1700, 4900, 2760])

    add_heading(doc, "10. Resultado esperado de la revisión")
    add_para(doc, "Al finalizar deben existir 12 decisiones completas, fechadas y atribuibles. El resultado no es una redacción ideal ni un diagnóstico: es un conjunto de reglas de aceptación para evaluar de manera consistente las futuras decisiones de Sol.")
    add_callout(doc, "Siguiente paso.", "Cargar el gold aprobado desde la consola interna de Curación Tier 1. El sistema validará que no falte ningún campo antes de permitir la calibración. Si algo es incierto, dejarlo explícito como limitación o elegir withheld.", fill=CAUTION)

    return doc


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--gold-manifest", required=True)
    parser.add_argument("--review-cards", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    gold_manifest = json.loads(Path(args.gold_manifest).read_text(encoding="utf-8"))
    review_cards = json.loads(Path(args.review_cards).read_text(encoding="utf-8"))
    document = build_document(gold_manifest, review_cards)
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = "Guía de revisión - Recuración Tier 1"
    document.core_properties.subject = "Guía práctica del gold standard interno"
    document.core_properties.author = "HEAL by FON"
    document.save(output)
    print(json.dumps({"output": str(output.resolve()), "status": "created"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
