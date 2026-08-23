#!/usr/bin/env python3
"""Create human and technical audit packages for the 12+93 Tier 1 candidate."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import subprocess
import zipfile
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


EXPECTED_REGISTRY = "73b94c09c31c184135bc16b2e756fa447f4c040a8bf38b49181168a8c62a967f"


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def flatten(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, list):
        return " | ".join(flatten(item) for item in value)
    if isinstance(value, dict):
        return " | ".join(f"{key}: {flatten(item)}" for key, item in value.items())
    return str(value)


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    headers: list[str] = []
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)
    if not headers:
        headers, rows = ["estado"], [{"estado": "Sin datos"}]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows({key: flatten(row.get(key)) for key in headers} for row in rows)


def collect(campaign: Path, registry: Path) -> dict[str, Any]:
    summary = read_json(campaign / "candidate_summary.json")
    if (campaign / "terminal.json").exists():
        terminal = read_json(campaign / "terminal.json")
    else:
        terminal = summary
    candidates = read_jsonl(campaign / "tier1_prototype_snapshot_candidate.jsonl")
    records = [read_json(path) for path in sorted((campaign / "records").glob("*.json"))]
    record_by_group = {row["group_id"]: row for row in records}
    decisions: list[dict[str, Any]] = []
    roles: list[dict[str, Any]] = []
    sources: list[dict[str, Any]] = []
    calls: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    arbiters: list[dict[str, Any]] = []
    for candidate in sorted(candidates, key=lambda row: row["group_id"]):
        group_id = candidate["group_id"]
        record = record_by_group.get(group_id, {})
        direction = candidate.get("direction_assessment") or {}
        decisions.append({
            "grupo": group_id,
            "módulo": group_id.split(":", 1)[1] if ":" in group_id else "",
            "estado": candidate.get("core_status"),
            "contexto_utilizable": candidate.get("context_usable"),
            "techo_inferencia": candidate.get("inference_ceiling"),
            "dirección": candidate.get("dominant_direction") or direction.get("dominant_direction"),
            "conflicto_material": candidate.get("material_conflict", direction.get("material_conflict")),
            "confianza": candidate.get("confidence"),
            "provenance_allowlist": candidate.get("allowlist_provenance"),
            "arbitraje": candidate.get("adjudicated", False),
            "limitaciones": candidate.get("limitations"),
            "evidence_ids": candidate.get("evidence_ids"),
            "packet_sha256": candidate.get("packet_sha256"),
        })
        if record.get("adjudicated"):
            arbiters.append({
                "grupo": group_id,
                "motivos": record.get("arbitration_reasons"),
                "estado_final": candidate.get("core_status"),
                "techo_final": candidate.get("inference_ceiling"),
            })
        for role in ("curator", "critic", "arbiter"):
            semantic = record.get(f"{role}_semantic")
            normalized = record.get(f"{role}_decision")
            role_errors = record.get(f"{role}_errors") or []
            if semantic:
                ndirection = (normalized or {}).get("direction_assessment") or {}
                roles.append({
                    "grupo": group_id, "rol": role,
                    "estado": (normalized or {}).get("core_status"),
                    "techo": (normalized or {}).get("inference_ceiling"),
                    "dirección": ndirection.get("dominant_direction"),
                    "conflicto_material": ndirection.get("material_conflict"),
                    "confianza": semantic.get("confidence"),
                    "errores": role_errors,
                })
                for item in semantic.get("used_evidence") or []:
                    sources.append({
                        "grupo": group_id, "rol": role, "evidence_id": item.get("evidence_id"),
                        "disposición": "usada", "roles": item.get("roles"),
                        "dirección": item.get("supports_direction"),
                        "clase_conflicto": item.get("core_conflict_class"),
                        "motivo": item.get("core_conflict_rationale"),
                    })
                for item in semantic.get("excluded_evidence") or []:
                    sources.append({
                        "grupo": group_id, "rol": role, "evidence_id": item.get("evidence_id"),
                        "disposición": "excluida", "roles": "", "dirección": "",
                        "clase_conflicto": "", "motivo": item.get("reason_code"),
                    })
            for error in role_errors:
                errors.append({"grupo": group_id, "rol": role, "error": error})
        for audit in record.get("call_audits") or []:
            for failure in audit.get("technical_failures") or []:
                errors.append({
                    "grupo": group_id,
                    "rol": audit.get("role"),
                    "error": failure,
                })
            for attempt in audit.get("attempts") or []:
                calls.append({
                    "grupo": group_id, "rol": audit.get("role"), "intento": attempt.get("attempt"),
                    "response_id": attempt.get("response_id"), "modelo": attempt.get("effective_model"),
                    "estado": attempt.get("response_status") or attempt.get("http_status"),
                    "input_tokens": attempt.get("input_tokens"),
                    "cached_input_tokens": attempt.get("cached_input_tokens"),
                    "output_tokens": attempt.get("output_tokens"),
                    "reasoning_tokens": attempt.get("reasoning_tokens"),
                    "latencia_segundos": attempt.get("latency_seconds"),
                    "costo_usd": attempt.get("estimated_cost_usd"),
                    "costo_observable": attempt.get("cost_observability"),
                })
    for error in summary.get("critical_errors") or []:
        errors.append({"grupo": str(error).split(":", 1)[0], "rol": "final", "error": error})
    if summary.get("structural_stop_reason"):
        errors.append({"grupo": "campaña", "rol": "structural_stop", "error": summary["structural_stop_reason"]})

    # Human review: every approval plus a deterministic 20% within each
    # status/module stratum for withheld/rejected decisions.
    selected: set[str] = {
        row["grupo"] for row in decisions if row["estado"] in {"approved", "approved_with_conflict"}
    }
    for status in ("withheld", "rejected"):
        modules = sorted({row["módulo"] for row in decisions if row["estado"] == status})
        for module in modules:
            stratum = sorted(
                (row for row in decisions if row["estado"] == status and row["módulo"] == module),
                key=lambda row: hashlib.sha256(row["grupo"].encode()).hexdigest(),
            )
            selected.update(row["grupo"] for row in stratum[: math.ceil(len(stratum) * 0.20)])
    review = []
    for row in decisions:
        if row["grupo"] not in selected:
            continue
        review.append({
            "revisión_requerida": "100% aprobaciones" if row["estado"] in {"approved", "approved_with_conflict"} else "muestra 20% estratificada",
            **row,
            "decisión_revisor": "",
            "comentarios_revisor": "",
            "responsable": "",
            "fecha": "",
        })
    estimate = read_json(campaign / "cost_estimate.json")
    ledger = read_json(campaign / "budget_ledger.json")
    actual_registry = sha256(registry)
    summary_rows = [
        {"indicador": "Estado del candidato", "valor": summary.get("status"), "detalle": "No publicado; requiere revisión humana"},
        {"indicador": "Grupos válidos", "valor": summary.get("total_groups_valid"), "detalle": "Meta: 105"},
        {"indicador": "Grupos firmados heredados", "valor": summary.get("signed_groups_carried"), "detalle": "Provenance signed_gold_prototype_replay"},
        {"indicador": "Grupos nuevos evaluados", "valor": summary.get("unsigned_groups_attempted"), "detalle": "Provenance packet_selected"},
        {"indicador": "Arbitrajes", "valor": summary.get("adjudications"), "detalle": "Sólo por desacuerdo científico o baja confianza"},
        {"indicador": "Costo estimado", "valor": summary.get("estimated_cost_usd"), "detalle": "USD con P95 y contingencia"},
        {"indicador": "Costo contabilizado", "valor": summary.get("accounted_cost_usd"), "detalle": "USD; hard cap 75"},
        {"indicador": "Grupos incompletos", "valor": len(summary.get("critical_errors") or []), "detalle": summary.get("structural_stop_reason") or "Sin stop estructural"},
        {"indicador": "Registry activo", "valor": "INTACTO" if actual_registry == EXPECTED_REGISTRY else "CAMBIÓ", "detalle": actual_registry},
        {"indicador": "Próximo gate", "valor": summary.get("next_gate"), "detalle": "No publicar antes de la revisión"},
    ]
    return {
        "summary": summary, "terminal": terminal, "summary_rows": summary_rows,
        "decisions": decisions, "roles": roles, "sources": sources, "calls": calls,
        "errors": errors, "arbiters": arbiters, "review": review,
        "estimate": estimate, "ledger": ledger, "registry_sha256": actual_registry,
    }


def add_doc_table(doc: Document, headers: list[str], rows: list[list[Any]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
        for run in table.rows[0].cells[index].paragraphs[0].runs:
            run.bold = True; run.font.size = Pt(8)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = flatten(value)
            for run in cells[index].paragraphs[0].runs:
                run.font.size = Pt(7.5)


def create_docx(path: Path, data: dict[str, Any]) -> None:
    doc = Document(); section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[style_name].font.name = "Calibri"
        doc.styles[style_name].font.color.rgb = RGBColor(31, 77, 120)
    title = doc.add_paragraph(style="Title"); title.add_run("HEAL — candidato Tier 1 (12 + 93)")
    subtitle = doc.add_paragraph("Auditoría interna previa a cualquier publicación")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("Resultado ejecutivo", 1)
    summary = data["summary"]
    doc.add_paragraph(
        f"Estado: {summary.get('status')}. Se obtuvieron {summary.get('total_groups_valid')} decisiones válidas "
        f"de 105 y se contabilizaron USD {float(summary.get('accounted_cost_usd') or 0):.4f}. "
        "Este snapshot no está publicado y no constituye validación clínica independiente."
    )
    add_doc_table(doc, ["Indicador", "Valor", "Detalle"], [
        [row["indicador"], row["valor"], row["detalle"]] for row in data["summary_rows"]
    ])
    doc.add_heading("Distribución de decisiones", 1)
    add_doc_table(doc, ["Estado", "Cantidad"], [[key, value] for key, value in sorted((summary.get("status_counts") or {}).items())])
    doc.add_heading("Techos de inferencia", 1)
    add_doc_table(doc, ["Techo", "Cantidad"], [[key, value] for key, value in sorted((summary.get("inference_ceiling_counts") or {}).items())])
    doc.add_heading("Cómo revisar", 1)
    for text in (
        "Revisar en el Excel el 100% de approved y approved_with_conflict.",
        "Revisar la muestra estratificada del 20% de withheld y rejected.",
        "Confirmar que packet_selected no se confunda con evidencia firmada por una persona.",
        "No publicar ni reemplazar el registry activo hasta cerrar la revisión humana.",
    ):
        doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("Controles", 1)
    doc.add_paragraph(f"Registry activo: {data['registry_sha256']}")
    doc.add_paragraph(f"Arbitrajes: {summary.get('adjudications')}; errores listados: {len(data['errors'])}.")
    doc.add_paragraph("Luna, VCF y publicación del snapshot: fuera de esta campaña.")
    doc.save(path)


def create_pdf(path: Path, data: dict[str, Any]) -> None:
    styles = getSampleStyleSheet()
    title = ParagraphStyle("HealTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=colors.HexColor("#0B2545"))
    heading = ParagraphStyle("HealHeading", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=13, textColor=colors.HexColor("#2E74B5"), spaceBefore=10, spaceAfter=6)
    body = ParagraphStyle("HealBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=12, spaceAfter=6)
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=16*mm, rightMargin=16*mm, topMargin=16*mm, bottomMargin=16*mm)
    summary = data["summary"]
    story: list[Any] = [Paragraph("HEAL — candidato Tier 1 (12 + 93)", title), Paragraph("Auditoría interna previa a cualquier publicación", body), Spacer(1, 5)]
    story += [Paragraph("Resultado ejecutivo", heading), Paragraph(
        f"Estado: {summary.get('status')}. Decisiones válidas: {summary.get('total_groups_valid')}/105. "
        f"Costo contabilizado: USD {float(summary.get('accounted_cost_usd') or 0):.4f}. "
        "El candidato permanece bloqueado hasta revisión humana y no constituye validación clínica.", body)]
    rows = [["Indicador", "Valor"]] + [[row["indicador"], flatten(row["valor"])] for row in data["summary_rows"]]
    table = Table(rows, colWidths=[70*mm, 100*mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#E8EEF5")), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTNAME", (0,1), (-1,-1), "Helvetica"), ("FONTSIZE", (0,0), (-1,-1), 8),
        ("GRID", (0,0), (-1,-1), 0.35, colors.HexColor("#CBD5E1")), ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ])); story.append(table)
    story += [Paragraph("Distribución", heading)]
    distribution = "Estados: " + ", ".join(f"{key}={value}" for key, value in sorted((summary.get("status_counts") or {}).items()))
    ceilings = "Techos: " + ", ".join(f"{key}={value}" for key, value in sorted((summary.get("inference_ceiling_counts") or {}).items()))
    story += [Paragraph(distribution, body), Paragraph(ceilings, body), Paragraph("Revisión requerida", heading), Paragraph(
        "El Excel contiene el 100% de approved/approved_with_conflict y una muestra estratificada del 20% de withheld/rejected. "
        "El registry activo no fue modificado; Luna y VCF no se ejecutaron en esta campaña.", body)]
    doc.build(story)


def scan_no_secrets(paths: list[Path]) -> None:
    secret = os.environ.get("HEAL_OPENAI_API_KEY", "").encode()
    for path in paths:
        data = path.read_bytes()
        lowered = data.lower()
        if secret and len(secret) >= 16 and secret in data:
            raise RuntimeError(f"Secret found in {path}")
        if b"authorization: bearer" in lowered:
            raise RuntimeError(f"Authorization header found in {path}")


def zip_tree(target: Path, roots: list[tuple[Path, str]], *, no_json: bool) -> None:
    files: list[tuple[Path, str]] = []
    for root, prefix in roots:
        for path in sorted(root.rglob("*")):
            if not path.is_file() or path == target:
                continue
            if no_json and path.suffix.lower() in {".json", ".jsonl", ".ndjson"}:
                continue
            files.append((path, f"{prefix}/{path.relative_to(root).as_posix()}".lstrip("/")))
    scan_no_secrets([path for path, _ in files])
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path, name in files:
            archive.write(path, name)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--campaign", required=True)
    parser.add_argument("--evidence-manifest-dir", required=True)
    parser.add_argument("--registry", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--node")
    parser.add_argument("--node-path")
    parser.add_argument("--workbook-builder")
    args = parser.parse_args()
    campaign = Path(args.campaign).resolve(); registry = Path(args.registry).resolve()
    output = Path(args.output_dir).resolve(); human = output / "Paquete_humano"
    technical = output / "Auditoria_tecnica"; qa = output / "QA_XLSX"
    human.mkdir(parents=True, exist_ok=True); technical.mkdir(parents=True, exist_ok=True)
    data = collect(campaign, registry)
    if data["registry_sha256"] != EXPECTED_REGISTRY:
        raise RuntimeError("Protected active registry changed")
    csv_sets = {
        "RESUMEN.csv": data["summary_rows"], "DECISIONES_105.csv": data["decisions"],
        "ROLES_CURATOR_CRITIC_ARBITER.csv": data["roles"], "FUENTES_93.csv": data["sources"],
        "TELEMETRIA_COSTOS.csv": data["calls"], "ARBITRAJES.csv": data["arbiters"],
        "ERRORES_CUARENTENA.csv": data["errors"], "REVISION_HUMANA.csv": data["review"],
    }
    for name, rows in csv_sets.items():
        write_csv(human / name, rows)
    create_docx(human / "Informe_Candidato_Tier1_105.docx", data)
    create_pdf(human / "Informe_Candidato_Tier1_105.pdf", data)
    (human / "GUIA_DE_LECTURA.md").write_text(
        "# Guía de lectura\n\n1. Abra el PDF para el resumen ejecutivo.\n"
        "2. Abra el Excel y empiece por RESUMEN.\n3. Use REVISION_HUMANA para firmar o corregir.\n"
        "4. DECISIONES_105 contiene el snapshot completo; FUENTES_93 conserva la trazabilidad de Sol.\n"
        "5. packet_selected es selección automática, no firma humana.\n\n"
        "No publicar este candidato antes de completar la revisión indicada.\n", encoding="utf-8"
    )
    workbook = {
        "RESUMEN": data["summary_rows"], "DECISIONES_105": data["decisions"],
        "REVISION_HUMANA": data["review"], "ROLES": data["roles"],
        "ARBITRAJES": data["arbiters"], "ERRORES": data["errors"],
        "FUENTES": data["sources"], "TELEMETRIA_COSTOS": data["calls"],
    }
    workbook_data = technical / "workbook_data.json"
    workbook_data.write_text(json.dumps(workbook, ensure_ascii=False), encoding="utf-8")
    if args.node and args.workbook_builder:
        env = dict(os.environ)
        if args.node_path: env["NODE_PATH"] = args.node_path
        subprocess.run([
            args.node, args.workbook_builder, str(workbook_data),
            str(human / "HEAL_Candidato_Tier1_105_Auditoria.xlsx"), str(qa),
        ], check=True, env=env)
    else:
        raise RuntimeError("Node runtime and workbook builder are required")
    shutil.copy2(campaign / "candidate_summary.json", technical / "candidate_summary.json")
    shutil.copy2(campaign / "cost_estimate.json", technical / "cost_estimate.json")
    shutil.copy2(campaign / "budget_ledger.json", technical / "budget_ledger.json")
    for inspect_file in human.glob("*.inspect.ndjson"):
        inspect_file.unlink()
    manifest_rows = []
    for path in sorted(human.rglob("*")):
        if path.is_file() and path.name != "MANIFEST_SHA256.csv":
            manifest_rows.append({"archivo": path.relative_to(human).as_posix(), "bytes": path.stat().st_size, "sha256": sha256(path)})
    write_csv(human / "MANIFEST_SHA256.csv", manifest_rows)
    zip_tree(output / "HEAL_Candidato_Tier1_105_HUMANO.zip", [(human, "")], no_json=True)
    zip_tree(output / "HEAL_Candidato_Tier1_105_TECNICO.zip", [
        (campaign, "campaign"), (Path(args.evidence_manifest_dir).resolve(), "evidence_manifest"),
        (technical, "package_support"),
    ], no_json=False)
    print(json.dumps({"status": "tier1_candidate_audit_packaged", "output": str(output)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
