#!/usr/bin/env python3
"""Build the human-readable audit package for a stopped Sol candidate-1 campaign."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_csv(path: Path, rows: list[dict], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def join(values) -> str:
    return "; ".join(str(value) for value in (values or []))


def set_cell_fill(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    props.append(fill)


def set_repeat_table_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    props.append(header)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_fill(cell, "E8EEF5")
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
    set_repeat_table_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
    if widths:
        dxa_widths = [round(width * 1440) for width in widths]
        table_width = sum(dxa_widths)
        table_props = table._tbl.tblPr
        tbl_width = table_props.first_child_found_in("w:tblW")
        if tbl_width is None:
            tbl_width = OxmlElement("w:tblW")
            table_props.append(tbl_width)
        tbl_width.set(qn("w:type"), "dxa")
        tbl_width.set(qn("w:w"), str(table_width))
        tbl_indent = OxmlElement("w:tblInd")
        tbl_indent.set(qn("w:type"), "dxa")
        tbl_indent.set(qn("w:w"), "120")
        table_props.append(tbl_indent)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        table_props.append(layout)
        grid_columns = list(table._tbl.tblGrid.gridCol_lst)
        for index, width in enumerate(dxa_widths):
            grid_columns[index].set(qn("w:w"), str(width))
        for row in table.rows:
            for index, width in enumerate(dxa_widths):
                cell = row.cells[index]
                cell.width = Inches(width / 1440)
                cell_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
                cell_width.set(qn("w:type"), "dxa")
                cell_width.set(qn("w:w"), str(width))
    return table


def build_docx(output: Path, terminal: dict, probe: dict, gold: dict, price: dict, source_url: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color in (("Title", 24, "0B2545"), ("Heading 1", 16, "2E74B5"), ("Heading 2", 13, "2E74B5")):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    header = section.header.paragraphs[0]
    header.text = "HEAL | Auditoria Tier 1 - candidate-1"
    header.style = doc.styles["Caption"]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Uso interno - 16/08/2026")

    doc.add_heading("Evaluacion Sol candidate-1", 0)
    subtitle = doc.add_paragraph("Informe de ejecucion y trazabilidad - Gold Tier 1")
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.color.rgb = RGBColor.from_string("4B5563")

    status = doc.add_paragraph()
    status_run = status.add_run("RESULTADO: DETENIDO ANTES DE CALIBRATION")
    status_run.bold = True
    status_run.font.color.rgb = RGBColor.from_string("9B1C1C")
    doc.add_paragraph(
        "La prueba minima llego a la API de OpenAI, pero los dos intentos tecnicos permitidos fueron rechazados "
        "por cuota insuficiente (HTTP 429). No se ejecuto ningun grupo cientifico, no se optimizo el prompt y no se inicio holdout."
    )

    doc.add_heading("1. Resumen ejecutivo", level=1)
    add_table(doc, ["Control", "Resultado"], [
        ["Gold v4", "Valido: 12 grupos firmados y 0 errores de validate-gold"],
        ["Preflight Sol", "Fallido: insufficient_quota en 2/2 intentos"],
        ["Calibration", "No iniciada (0/6 grupos)"],
        ["Optimizacion", "No ejecutada"],
        ["Prompt congelado", "No existe"],
        ["Holdout", "No iniciado (0/6 grupos)"],
        ["Registro activo", "Sin cambios; hash antes y despues coincidente"],
        ["Etapas fuera de alcance", "105 grupos, Luna y VCF: no ejecutadas"],
    ], [1.75, 4.75])

    doc.add_heading("2. Que significa este resultado", level=1)
    doc.add_paragraph(
        "Este resultado no permite evaluar la calidad cientifica de Sol. Es un bloqueo operativo de la cuenta/proyecto asociado a la API key. "
        "Tampoco es un fallo del Gold, de los packets ni del prompt candidate-1."
    )
    for text in (
        "No hay salidas curator, critic ni arbiter que comparar con el Gold.",
        "No se consumieron tokens observables: la API no devolvio una Response exitosa con telemetria.",
        "El costo registrado es no observable, no cero: ante un rechazo previo a la respuesta no hay usage reportado.",
        "El protocolo impidio reintentos adicionales y dejo un estado terminal auditable.",
    ):
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("3. Separacion de fases", level=1)
    phase_rows = [
        ["Preflight", "Ejecutado", "2 intentos tecnicos", "Detenido por cuota"],
        ["Calibration", "No ejecutado", "0/6", "Sin feedback semantico"],
        ["Optimize", "No ejecutado", "0 candidatos adicionales", "Holdout no expuesto"],
        ["Freeze prompt", "No ejecutado", "Sin manifest", "No corresponde"],
        ["Holdout", "No ejecutado", "0/6", "Sigue sin utilizar"],
    ]
    add_table(doc, ["Fase", "Estado", "Cobertura", "Observacion"], phase_rows, [1.25, 1.35, 1.45, 2.45])

    doc.add_heading("4. Gold firmado disponible", level=1)
    cards = gold.get("cards") or []
    rows = [[
        card.get("group_id", ""), card.get("split", ""), join(card.get("acceptable_core_statuses")),
        str(card.get("context_usable")), join(card.get("acceptable_inference_ceilings")),
    ] for card in cards]
    add_table(doc, ["Grupo", "Split", "Estado aceptable", "Contexto", "Techo"], rows, [1.1, 0.85, 1.75, 0.8, 2.0])
    doc.add_paragraph(
        "La hoja GOLD_VS_SOL del Excel conserva estos 12 objetivos y marca todas las columnas Sol como no ejecutadas."
    )

    doc.add_heading("5. Telemetria y costo", level=1)
    attempts = probe.get("attempts") or []
    attempt_rows = [[
        item.get("attempt", ""), item.get("http_response_received", False), item.get("effective_model", "") or "No informado",
        item.get("input_tokens", 0), item.get("output_tokens", 0),
        "No observable" if item.get("estimated_cost_usd") is None else f"USD {item['estimated_cost_usd']:.8f}",
    ] for item in attempts]
    add_table(doc, ["Intento", "Response", "Modelo efectivo", "Input", "Output", "Costo"], attempt_rows, [0.65, 0.75, 1.6, 0.8, 0.8, 1.9])
    doc.add_paragraph(
        f"Tarifa congelada para calculos futuros: USD {price['input_per_million']}/M input no cacheado, "
        f"USD {price['cached_input_per_million']}/M input cacheado y USD {price['output_per_million']}/M output. "
        f"Fuente: {source_url}"
    )

    doc.add_heading("6. Integridad y alcance", level=1)
    add_table(doc, ["Elemento", "Valor"], [
        ["Hash registro antes", terminal.get("registry_sha256_before", "")],
        ["Hash registro despues", terminal.get("registry_sha256_after", "")],
        ["Gold manifest", gold.get("manifest_sha256", "")],
        ["Estado terminal", terminal.get("status", "")],
    ], [1.65, 4.85])
    doc.add_paragraph(
        "La igualdad del hash del registro confirma que la ejecucion no publico ni modifico el canon activo. "
        "La carpeta tecnica conserva los estados JSON del probe y del protocolo, junto con los hashes de los artefactos."
    )

    doc.add_page_break()
    doc.add_heading("7. Proximos pasos y limitaciones", level=1)
    steps = [
        "Habilitar billing/cuota en el proyecto asociado a HEAL_OPENAI_API_KEY o reemplazarla por una clave de un proyecto con cuota para gpt-5.6-sol.",
        "Confirmar con una prueba minima nueva en una campaña nueva. No reutilizar los dos intentos agotados de este run.",
        "Si el nuevo preflight pasa, ejecutar el mismo orquestador: calibration, optimize solo con calibration, freeze-prompt y holdout una vez.",
        "Generar el paquete completo con salidas por grupo. Este informe debe conservarse como antecedente operativo, no como evaluacion cientifica de Sol.",
    ]
    for step in steps:
        doc.add_paragraph(step, style="List Number")

    limitations = doc.add_paragraph()
    limitations.paragraph_format.space_before = Pt(6)
    limitations.add_run("Limitaciones. ").bold = True
    limitations.add_run(
        "No hay conclusion sobre precision, consistencia, conflictos, evidencia utilizada o concordancia con el Gold porque no se proceso ningun grupo. "
        "El holdout no fue visto por el modelo y permanece metodologicamente disponible, pero para cumplir la regla de intentos debera utilizarse en una nueva ejecucion logica."
    )

    doc.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--campaign-root", required=True)
    parser.add_argument("--evidence-manifest", required=True)
    parser.add_argument("--gold", required=True)
    parser.add_argument("--registry", required=True)
    parser.add_argument("--service-root", required=True)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()

    campaign = Path(args.campaign_root).resolve()
    evidence_path = Path(args.evidence_manifest).resolve()
    gold_path = Path(args.gold).resolve()
    registry_path = Path(args.registry).resolve()
    service = Path(args.service_root).resolve()
    output = Path(args.output_dir).resolve()
    output.mkdir(parents=True, exist_ok=True)

    terminal = read_json(campaign / "protocol_terminal_state.json")
    probe = read_json(campaign / "preflight" / "probe_audit.json")
    price = read_json(campaign / "price_snapshot.json")
    gold = read_json(gold_path)
    evidence = read_json(evidence_path)
    packets = {row["group_id"]: read_json(Path(row["packet_path"])) for row in evidence["packets"]}

    comparisons = []
    for card in gold["cards"]:
        comparisons.append({
            "fase": card["split"], "grupo": card["group_id"],
            "gold_core_status": join(card["acceptable_core_statuses"]),
            "gold_context_usable": card["context_usable"],
            "gold_inference_ceiling": join(card["acceptable_inference_ceilings"]),
            "gold_directions": join(card["acceptable_directions"]),
            "gold_valid_evidence_ids": join(card["valid_evidence_ids"]),
            "gold_invalid_evidence_ids": join(card["invalid_evidence_ids"]),
            "sol_curator": "NO EJECUTADO", "sol_critic": "NO EJECUTADO", "sol_arbiter": "NO EJECUTADO",
            "sol_final": "NO EJECUTADO", "comparacion": "NO EVALUABLE - PREFLIGHT DE CUOTA FALLIDO",
        })

    sources = []
    gold_by_group = {card["group_id"]: card for card in gold["cards"]}
    for group_id, packet in packets.items():
        card = gold_by_group[group_id]
        selected = set(packet.get("selected_evidence_ids") or [])
        valid = set(card.get("valid_evidence_ids") or [])
        invalid = set(card.get("invalid_evidence_ids") or [])
        for source in packet.get("source_ledger") or []:
            evidence_id = source.get("evidence_id", "")
            sources.append({
                "grupo": group_id, "evidence_id": evidence_id, "seleccionada_packet": evidence_id in selected,
                "valida_firmada": evidence_id in valid, "invalida_firmada": evidence_id in invalid,
                "tipo": source.get("source_kind", ""), "titulo": source.get("title", ""),
                "pmid": source.get("pmid", ""), "doi": source.get("doi", ""),
                "fecha_publicacion": source.get("publication_date", ""), "estado_elegibilidad": source.get("eligibility_status", ""),
                "motivos_exclusion": join(source.get("exclusion_reasons")), "url": source.get("url", ""),
            })

    telemetry = []
    failures = probe.get("technical_failures") or []
    for item in probe.get("attempts") or []:
        index = int(item.get("attempt") or 0)
        telemetry.append({
            "fase": "preflight", "rol": "probe", "intento": index, "modelo_solicitado": probe.get("model", ""),
            "modelo_efectivo": item.get("effective_model", ""), "response_id": item.get("response_id", ""),
            "http_response_recibida": item.get("http_response_received", False), "latencia_segundos": item.get("latency_seconds", 0),
            "input_tokens": item.get("input_tokens", 0), "cached_input_tokens": item.get("cached_input_tokens", 0),
            "uncached_input_tokens": item.get("uncached_input_tokens", 0), "output_tokens": item.get("output_tokens", 0),
            "total_tokens": item.get("total_tokens", 0), "costo_estimado_usd": item.get("estimated_cost_usd"),
            "observabilidad_costo": item.get("cost_observability", ""),
            "error": "HTTP 429 insufficient_quota - revisar cuota y billing del proyecto API"
            if 0 < index <= len(failures) and "insufficient_quota" in failures[index - 1]
            else (failures[index - 1] if 0 < index <= len(failures) else ""),
        })

    prompts = []
    for role in ("curator", "critic", "arbiter", "optimizer"):
        path = service / f"prompt_mechanism_{role}_v2.md" if role != "optimizer" else service / "prompt_optimizer_v2.md"
        prompts.append({
            "candidate": "candidate-1", "rol": role, "ruta": str(path), "sha256": sha256_file(path),
            "estado": "INICIAL; NO EJECUTADO" if role != "optimizer" else "NO EJECUTADO",
            "motivo_cambio": "Sin cambios; calibration no inicio",
        })

    phases = [
        {"fase": "preflight", "estado": "FALLIDO", "grupos_previstos": 0, "grupos_ejecutados": 0, "detalle": "2/2 intentos: HTTP 429 insufficient_quota"},
        {"fase": "calibration", "estado": "NO INICIADA", "grupos_previstos": 6, "grupos_ejecutados": 0, "detalle": "Sin salidas Sol"},
        {"fase": "optimize", "estado": "NO EJECUTADA", "grupos_previstos": 0, "grupos_ejecutados": 0, "detalle": "Sin feedback semantico"},
        {"fase": "freeze-prompt", "estado": "NO EJECUTADA", "grupos_previstos": 0, "grupos_ejecutados": 0, "detalle": "No hay prompt congelado"},
        {"fase": "holdout", "estado": "NO INICIADO", "grupos_previstos": 6, "grupos_ejecutados": 0, "detalle": "Holdout no expuesto"},
    ]

    errors = [
        {"fase": "preflight", "codigo": "invalid_json_schema", "severidad": "implementacion_corregida", "intentos": 2,
         "detalle": "Probe inicial rechazado por faltar type en el campo const; artefacto preservado como probe_audit_failed_schema_v0.json."},
        {"fase": "preflight", "codigo": "insufficient_quota", "severidad": "bloqueante", "intentos": 2,
         "detalle": "La cuenta/proyecto de la API key no dispone de cuota. Calibration y holdout no se iniciaron."},
    ]

    hash_paths = [gold_path, evidence_path, registry_path, campaign / "price_snapshot.json", campaign / "preflight" / "probe_audit.json",
                  campaign / "preflight" / "probe_audit_failed_schema_v0.json", campaign / "protocol_terminal_state.json",
                  service / "mechanism_curation_v2.schema.json", service / "run_curation_v2.py", service / "curation_v2.py"]
    hash_paths.extend(Path(row["packet_path"]) for row in evidence["packets"])
    hash_paths.extend(service / name for name in ("prompt_mechanism_curator_v2.md", "prompt_mechanism_critic_v2.md", "prompt_mechanism_arbiter_v2.md", "prompt_optimizer_v2.md"))
    hashes = [{"archivo": str(path), "sha256": sha256_file(path), "tamano_bytes": path.stat().st_size} for path in hash_paths if path.exists()]

    rows_by_sheet = {
        "RESUMEN": [{"campo": "Estado", "valor": terminal["status"]}, {"campo": "Calibration", "valor": "0/6"},
                    {"campo": "Holdout", "valor": "0/6"}, {"campo": "Causa", "valor": "HTTP 429 insufficient_quota"},
                    {"campo": "Registro activo intacto", "valor": terminal["registry_sha256_before"] == terminal["registry_sha256_after"]},
                    {"campo": "Conclusion cientifica Sol", "valor": "No evaluable"}],
        "CALIBRATION": [row for row in phases if row["fase"] in {"preflight", "calibration", "optimize"}],
        "HOLDOUT": [row for row in phases if row["fase"] in {"freeze-prompt", "holdout"}],
        "GOLD_VS_SOL": comparisons,
        "CURATOR_CRITIC_ARBITER": [{"grupo": card["group_id"], "split": card["split"], "curator": "NO EJECUTADO", "critic": "NO EJECUTADO", "arbiter": "NO EJECUTADO", "final": "NO EJECUTADO"} for card in gold["cards"]],
        "PROMPTS": prompts,
        "ERRORES_CONFLICTOS": errors,
        "FUENTES": sources,
        "TELEMETRIA_COSTOS": telemetry,
        "HASHES": hashes,
    }

    write_csv(output / "comparacion_gold_vs_sol.csv", comparisons, list(comparisons[0]))
    write_csv(output / "fuentes_utilizadas_y_descartadas.csv", sources, list(sources[0]))
    write_csv(output / "telemetria_costos.csv", telemetry, list(telemetry[0]))
    write_csv(output / "errores_conflictos.csv", errors, list(errors[0]))
    write_csv(output / "hashes_sha256.csv", hashes, list(hashes[0]))
    (output / "workbook_data.json").write_text(json.dumps(rows_by_sheet, ensure_ascii=False, indent=2), encoding="utf-8")
    build_docx(output / "LLM1_Tier1_Sol_candidate-1_Informe_Ejecucion.docx", terminal, probe, gold, price, price["source_url"])
    print(json.dumps({"status": "built", "output": str(output), "rows_by_sheet": {key: len(value) for key, value in rows_by_sheet.items()}}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
