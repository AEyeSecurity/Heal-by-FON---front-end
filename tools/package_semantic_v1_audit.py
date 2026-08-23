from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import shutil
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


CAMPAIGN = Path(
    r"F:\Heal by FON\data\curation-candidates\tier1-human-review-20260810\reconciliation-v4"
    r"\sol-semantic-v1-verification-1-20260822"
)
RECONCILIATION = CAMPAIGN.parent
REPO = Path(r"F:\Heal by FON\app")
SERVICE = REPO / "services" / "heal-tier1-curation-v2"
GOLD_PATH = RECONCILIATION / "gold-v4" / "gold_frozen.json"
EVIDENCE_MANIFEST_PATH = RECONCILIATION / "evidence" / "evidence_manifest.json"
PACKETS_DIR = RECONCILIATION / "evidence" / "packets"
OUTPUT = Path(r"C:\Users\Usuario\Downloads\LLM1_Tier1_Sol_SemanticV1_2026-08-22")
ACTIVE_REGISTRY = Path(
    r"F:\Heal by FON\data\canon\curation\mechanism_registry_v1.csv"
)
EXPECTED_REGISTRY_SHA256 = "73b94c09c31c184135bc16b2e756fa447f4c040a8bf38b49181168a8c62a967f"
GROUP_ORDER = ["MTHFR:T1.1", "ASMT:T1.2", "FADS1:T1.3", "IL10:T1.4", "RUNX2:T1.5", "GCLM:T1.6"]
ROLE_ORDER = {"curator": 0, "critic": 1, "arbiter": 2}


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def dump_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def flatten_text(value: Any, separator: str = " | ") -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, list):
        return separator.join(flatten_text(item, separator) for item in value)
    if isinstance(value, dict):
        return separator.join(f"{key}: {flatten_text(item, separator)}" for key, item in value.items())
    return str(value)


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    headers: list[str] = []
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)
    if not headers:
        headers = ["estado"]
        rows = [{"estado": "Sin datos"}]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: flatten_text(row.get(key)) for key in headers})


def packet_path(group_id: str) -> Path:
    return PACKETS_DIR / (group_id.replace(":", "__") + ".json")


def role_file(directory: str, group_id: str, role: str) -> Path:
    return CAMPAIGN / "calibration" / directory / f"{group_id.replace(':', '__')}__{role}.json"


def used_ids(semantic: dict[str, Any] | None) -> set[str]:
    if not semantic:
        return set()
    return {item.get("evidence_id", "") for item in semantic.get("used_evidence", []) if item.get("evidence_id")}


def direction(decision: dict[str, Any] | None) -> dict[str, Any]:
    return (decision or {}).get("direction_assessment") or {}


def build_data() -> dict[str, Any]:
    terminal = read_json(CAMPAIGN / "protocol_terminal_state.json")
    config = read_json(CAMPAIGN / "protocol_configuration.json")
    phase = read_json(CAMPAIGN / "calibration" / "phase_acceptance.json")
    budget = read_json(CAMPAIGN / "budget_ledger.json")
    gold = read_json(GOLD_PATH)
    cards = {card["group_id"]: card for card in gold["cards"]}

    records: dict[str, dict[str, Any]] = {}
    roles: list[dict[str, Any]] = []
    comparisons: list[dict[str, Any]] = []
    sources: list[dict[str, Any]] = []
    telemetry: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []

    for audit_path in sorted((CAMPAIGN / "calibration" / "audit").glob("*.json")):
        audit = read_json(audit_path)
        for attempt in audit.get("attempts", []):
            telemetry.append(
                {
                    "fase": audit.get("phase"),
                    "grupo": audit.get("group_id"),
                    "rol": audit.get("role"),
                    "intento": attempt.get("attempt"),
                    "estado": attempt.get("response_status"),
                    "response_id": attempt.get("response_id"),
                    "modelo_solicitado": audit.get("model_id"),
                    "modelo_efectivo": attempt.get("effective_model"),
                    "reasoning_effort": attempt.get("reasoning_effort"),
                    "input_tokens": attempt.get("input_tokens"),
                    "cached_input_tokens": attempt.get("cached_input_tokens"),
                    "uncached_input_tokens": attempt.get("uncached_input_tokens"),
                    "output_tokens": attempt.get("output_tokens"),
                    "reasoning_tokens": attempt.get("reasoning_tokens"),
                    "total_tokens": attempt.get("total_tokens"),
                    "latencia_segundos": attempt.get("latency_seconds"),
                    "costo_usd": attempt.get("estimated_cost_usd"),
                    "observabilidad_costo": attempt.get("cost_observability"),
                    "retry": "Sí" if (attempt.get("attempt") or 1) > 1 else "No",
                    "timeout": "Sí" if attempt.get("response_status") == "timeout" else "No",
                    "prompt_sha256": audit.get("prompt_sha256"),
                    "schema_semantico_sha256": audit.get("semantic_schema_sha256"),
                    "schema_final_sha256": audit.get("final_schema_sha256"),
                    "evidencia_sha256": audit.get("evidence_sha256"),
                }
            )
        for message in audit.get("errors", []):
            errors.append({"fase": "calibration", "grupo": audit.get("group_id"), "rol": audit.get("role"), "tipo": "rol", "codigo_detalle": message})
        for message in audit.get("technical_failures", []):
            errors.append({"fase": "calibration", "grupo": audit.get("group_id"), "rol": audit.get("role"), "tipo": "técnico", "codigo_detalle": flatten_text(message)})

    for group_id in GROUP_ORDER:
        record = read_json(CAMPAIGN / "calibration" / "records" / f"{group_id.replace(':', '__')}.json")
        records[group_id] = record
        card = cards[group_id]
        valid = set(card.get("valid_evidence_ids", []))
        invalid = set(card.get("invalid_evidence_ids", []))
        role_names = ["curator", "critic"] + (["arbiter"] if record.get("arbiter_semantic") else [])
        for role in role_names:
            semantic = record.get(f"{role}_semantic") or read_json(role_file("semantic", group_id, role))
            decision = record.get(f"{role}_decision") or read_json(role_file("decisions", group_id, role))
            used = used_ids(semantic)
            unauthorized = sorted(used - valid)
            invalid_used = sorted(used & invalid)
            dirn = direction(decision)
            roles.append(
                {
                    "fase": "calibration",
                    "grupo": group_id,
                    "rol": role,
                    "core_status": decision.get("core_status"),
                    "identidad_exacta": decision.get("identity_exact"),
                    "relacion_modulo_directa": decision.get("module_relation_direct"),
                    "contexto_utilizable": decision.get("context_usable"),
                    "techo_inferencia": decision.get("inference_ceiling"),
                    "direccion": dirn.get("dominant_direction"),
                    "conflicto_material": dirn.get("material_conflict"),
                    "support_score": dirn.get("support_score"),
                    "opposition_score": dirn.get("opposition_score"),
                    "margin": dirn.get("margin"),
                    "confianza": decision.get("confidence"),
                    "evidencias_usadas": len(used),
                    "fuera_allowlist_gold": "; ".join(unauthorized),
                    "evidencia_invalida_usada": "; ".join(invalid_used),
                    "limitaciones": flatten_text(decision.get("limitations", [])),
                    "errores_rol": flatten_text(record.get(f"{role}_errors", [])),
                }
            )
            comparisons.extend(
                [
                    {
                        "grupo": group_id,
                        "rol": role,
                        "campo": "core_status",
                        "gold_aceptable": flatten_text(card.get("acceptable_core_statuses")),
                        "sol": decision.get("core_status"),
                        "cumple": decision.get("core_status") in card.get("acceptable_core_statuses", []),
                    },
                    {
                        "grupo": group_id,
                        "rol": role,
                        "campo": "inference_ceiling",
                        "gold_aceptable": flatten_text(card.get("acceptable_inference_ceilings")),
                        "sol": decision.get("inference_ceiling"),
                        "cumple": decision.get("inference_ceiling") in card.get("acceptable_inference_ceilings", []),
                    },
                    {
                        "grupo": group_id,
                        "rol": role,
                        "campo": "dominant_direction",
                        "gold_aceptable": flatten_text(card.get("acceptable_directions")),
                        "sol": dirn.get("dominant_direction"),
                        "cumple": dirn.get("dominant_direction") in card.get("acceptable_directions", []),
                    },
                    {
                        "grupo": group_id,
                        "rol": role,
                        "campo": "evidence_allowlist",
                        "gold_aceptable": f"Sólo {len(valid)} evidence IDs firmados",
                        "sol": "; ".join(sorted(used)),
                        "cumple": not unauthorized and not invalid_used,
                    },
                    {
                        "grupo": group_id,
                        "rol": role,
                        "campo": "limitations",
                        "gold_aceptable": "No vacías y compatibles con la limitación firmada",
                        "sol": flatten_text(decision.get("limitations", [])),
                        "cumple": bool(decision.get("limitations")),
                    },
                ]
            )
            for item in semantic.get("used_evidence", []):
                evidence_id = item.get("evidence_id")
                sources.append(
                    {
                        "grupo": group_id,
                        "rol": role,
                        "estado": "usada",
                        "evidence_id": evidence_id,
                        "roles_cientificos": flatten_text(item.get("roles", [])),
                        "direccion_semantica": item.get("supports_direction"),
                        "aplicabilidad_humana": item.get("human_applicability"),
                        "compatibilidad_funcional": item.get("functional_compatibility"),
                        "calidad": flatten_text(item.get("quality", {})),
                        "en_allowlist_gold": evidence_id in valid,
                        "firmada_invalida": evidence_id in invalid,
                        "motivo_exclusion": "",
                    }
                )
            for item in semantic.get("excluded_evidence", []) or []:
                sources.append(
                    {
                        "grupo": group_id,
                        "rol": role,
                        "estado": "descartada",
                        "evidence_id": item.get("evidence_id"),
                        "roles_cientificos": "",
                        "direccion_semantica": "",
                        "aplicabilidad_humana": "",
                        "compatibilidad_funcional": "",
                        "calidad": "",
                        "en_allowlist_gold": item.get("evidence_id") in valid,
                        "firmada_invalida": item.get("evidence_id") in invalid,
                        "motivo_exclusion": item.get("reason_code") or item.get("reason") or "",
                    }
                )

    for message in phase.get("errors", []):
        parts = str(message).split(":", 2)
        errors.append(
            {
                "fase": "calibration",
                "grupo": ":".join(parts[:2]) if len(parts) >= 2 else "",
                "rol": "runner",
                "tipo": "gate estructural",
                "codigo_detalle": message,
            }
        )
    errors.append(
        {
            "fase": "calibration",
            "grupo": "FADS1:T1.3",
            "rol": "curator",
            "tipo": "gap de implementación detectado post-run",
            "codigo_detalle": "El control por rol aceptó IDs presentes en el packet pero ausentes de valid_evidence_ids del Gold; el árbitro/final sí quedó dentro del Gold. Debe aplicarse la allowlist firmada antes de una nueva calibration.",
        }
    )

    group_rows: list[dict[str, Any]] = []
    for group_id in GROUP_ORDER:
        record = records[group_id]
        cur = record.get("curator_decision") or {}
        crit = record.get("critic_decision") or {}
        arb = record.get("arbiter_decision") or {}
        final = record.get("final_decision") or {}
        group_rows.append(
            {
                "grupo": group_id,
                "estado_gold": flatten_text(cards[group_id].get("acceptable_core_statuses")),
                "curator": cur.get("core_status"),
                "critic": crit.get("core_status"),
                "acuerdo_base": record.get("base_pair_agreement"),
                "motivo_arbitraje": flatten_text(record.get("arbitration_reasons", [])),
                "arbiter": arb.get("core_status", "No ejecutado"),
                "decision_final": final.get("core_status", "No disponible"),
                "direccion_final": direction(final).get("dominant_direction", "No disponible"),
                "conflicto_curator": direction(cur).get("material_conflict"),
                "conflicto_critic": direction(crit).get("material_conflict"),
                "conflicto_arbiter": direction(arb).get("material_conflict", "No ejecutado"),
                "resultado": "Válido" if final else "Detenido por gate",
            }
        )

    prompts = []
    for role in ("curator", "critic", "arbiter"):
        path = SERVICE / f"prompt_mechanism_{role}_semantic_v1.md"
        prompts.append(
            {
                "fase": "calibration",
                "rol": role,
                "version": "semantic-v1-verification-1 / baseline candidate-3-semantic-policy",
                "archivo": str(path),
                "sha256": sha256(path),
                "congelado": "No; calibration falló antes de freeze",
                "cambios_durante_campaña": "Ninguno; optimizer deshabilitado",
            }
        )

    hashes: list[dict[str, Any]] = []
    key_paths = [
        GOLD_PATH,
        EVIDENCE_MANIFEST_PATH,
        SERVICE / "mechanism_curation_semantic_v1.schema.json",
        SERVICE / "mechanism_curation_v2.schema.json",
        SERVICE / "curation_v2.py",
        SERVICE / "run_semantic_verification_v1.py",
        ACTIVE_REGISTRY,
        CAMPAIGN / "protocol_configuration.json",
        CAMPAIGN / "protocol_terminal_state.json",
        CAMPAIGN / "calibration" / "phase_acceptance.json",
    ] + [SERVICE / f"prompt_mechanism_{role}_semantic_v1.md" for role in ("curator", "critic", "arbiter")]
    for path in key_paths:
        hashes.append({"artefacto": path.name, "ruta": str(path), "sha256": sha256(path), "tipo": "registro activo" if path == ACTIVE_REGISTRY else "entrada/ejecución"})

    summary = [
        {"indicador": "Estado final", "valor": terminal.get("status"), "lectura": "Calibration falló; no hubo freeze ni holdout."},
        {"indicador": "Grupos completos", "valor": f"{phase.get('complete_groups', 6)}/6", "lectura": "Curator y critic respondieron para los seis grupos."},
        {"indicador": "Acuerdo curator/critic", "valor": f"{phase.get('pair_agreement', 4)}/6", "lectura": "El gate exigía 5/6."},
        {"indicador": "Arbitrajes", "valor": phase.get("arbiter_count", 1), "lectura": "FADS1 fue adjudicado; GCLM habría requerido un segundo y activó el stop."},
        {"indicador": "Errores de contrato/evidencia reportados durante roles", "valor": 0, "lectura": "No hubo fallos técnicos ni de schema en las respuestas API."},
        {"indicador": "Aprobaciones sin respaldo", "valor": phase.get("unsupported_approvals", 0), "lectura": "Cero en el scoring final de la fase."},
        {"indicador": "Costo total contabilizado", "valor": f"USD {budget.get('accounted_campaign_cost_usd', 0):.6f}", "lectura": "Incluye probe y calibration; sin costos no observables."},
        {"indicador": "Retries / timeouts", "valor": "0 / 0", "lectura": "Todas las respuestas fueron observables y completadas."},
        {"indicador": "Holdout", "valor": "NO INICIADO", "lectura": "Continúa virgen y no fue expuesto al modelo."},
        {"indicador": "Prompt final congelado", "valor": "NO", "lectura": "Freeze estaba condicionado a calibration aprobada."},
        {"indicador": "Registro activo", "valor": sha256(ACTIVE_REGISTRY), "lectura": "Intacto y coincide con el hash esperado."},
    ]

    holdout = [
        {
            "estado": "NO INICIADO",
            "motivo": "La calibration obtuvo acuerdo 4/6; el protocolo exige al menos 5/6.",
            "lock_global": "No creado",
            "grupos_expuestos": 0,
            "repetición_autorizada": "No aplicable; el holdout original sigue virgen",
        }
    ]
    arbitration = []
    for group_id in GROUP_ORDER:
        record = records[group_id]
        if record.get("arbitration_reasons") or not record.get("final_decision"):
            arbitration.append(
                {
                    "grupo": group_id,
                    "acuerdo_base": record.get("base_pair_agreement"),
                    "motivos": flatten_text(record.get("arbitration_reasons", [])),
                    "arbiter_ejecutado": bool(record.get("arbiter_decision")),
                    "resultado_arbiter": (record.get("arbiter_decision") or {}).get("core_status", "No ejecutado"),
                    "conflicto_curator": direction(record.get("curator_decision")).get("material_conflict"),
                    "conflicto_critic": direction(record.get("critic_decision")).get("material_conflict"),
                    "conflicto_arbiter": direction(record.get("arbiter_decision")).get("material_conflict", "No disponible"),
                    "lectura": "Desacuerdo limitado a clasificación de conflicto material; no al estado, dirección ni techo.",
                }
            )

    return {
        "terminal": terminal,
        "config": config,
        "phase": phase,
        "budget": budget,
        "gold": gold,
        "records": records,
        "cards": cards,
        "workbook": {
            "RESUMEN": summary,
            "CALIBRATION": group_rows,
            "HOLDOUT": holdout,
            "GOLD_VS_SOL": comparisons,
            "ROLES": roles,
            "ARBITRAJES": arbitration,
            "ERRORES": errors,
            "FUENTES": sources,
            "TELEMETRIA_COSTOS": sorted(telemetry, key=lambda row: (row["fase"], GROUP_ORDER.index(row["grupo"]) if row["grupo"] in GROUP_ORDER else 99, ROLE_ORDER.get(row["rol"], 99))),
            "PROMPTS": prompts,
            "HASHES": hashes,
        },
    }


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_width(cell, width_inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_width = tc_pr.find(qn("w:tcW"))
    if tc_width is None:
        tc_width = OxmlElement("w:tcW")
        tc_pr.append(tc_width)
    tc_width.set(qn("w:type"), "dxa")
    tc_width.set(qn("w:w"), str(int(width_inches * 1440)))


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_table(doc: Document, headers: list[str], rows: Iterable[Iterable[Any]], widths: list[float] | None = None, font_size: float = 8.5) -> None:
    rows = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, header in enumerate(headers):
        cell = header_row.cells[index]
        cell.text = header
        set_cell_shading(cell, "0B2545")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(font_size)
        if widths:
            set_cell_width(cell, widths[index])
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = flatten_text(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if widths:
                set_cell_width(cells[index], widths[index])
            if row_index % 2:
                set_cell_shading(cells[index], "F4F6F9")
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def build_docx(data: dict[str, Any], output_path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, before, after, color in (
        ("Title", 23, 0, 12, "111827"),
        ("Heading 1", 16, 16, 8, "2E74B5"),
        ("Heading 2", 13, 12, 6, "2E74B5"),
        ("Heading 3", 11.5, 10, 4, "1F4D78"),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.text = "HEAL by FON  |  Auditoría Tier 1 · Semantic V1"
    header.runs[0].font.size = Pt(8.5)
    header.runs[0].font.color.rgb = RGBColor.from_string("5B6573")
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.add_run("Evaluación controlada de Semantic V1 con GPT-5.6 Sol")
    subtitle = doc.add_paragraph()
    subtitle.add_run("Informe de calibration — holdout no iniciado").bold = True
    subtitle.runs[0].font.color.rgb = RGBColor.from_string("9B1C1C")
    doc.add_paragraph("Campaña: semantic-v1-verification-1 · Fecha: 22/08/2026 · Uso interno")

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    metadata = [
        ("Resultado", "CALIBRATION FALLIDA — detención correcta antes del freeze"),
        ("Acuerdo base", "4/6; el protocolo exigía 5/6"),
        ("Costo contabilizado", f"USD {data['budget']['accounted_campaign_cost_usd']:.6f}"),
        ("Registro activo", "Intacto; SHA-256 esperado y observado coinciden"),
    ]
    for row, (label, value) in zip(table.rows, metadata):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_width(row.cells[0], 1.55)
        set_cell_width(row.cells[1], 5.25)
        set_cell_shading(row.cells[0], "E8EEF5")
        for run in row.cells[0].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string("0B2545")

    add_heading(doc, "Resumen ejecutivo", 1)
    doc.add_paragraph(
        "La campaña confirmó que la nueva separación entre evaluación semántica de Sol y cálculo determinístico funciona end-to-end: las trece llamadas científicas completaron el schema, no hubo timeouts ni retries y desaparecieron los antiguos errores aritméticos direction_score_mismatch. Sin embargo, curator y critic coincidieron plenamente sólo en cuatro de los seis grupos."
    )
    doc.add_paragraph(
        "Las dos divergencias no cuestionan el mecanismo central: FADS1 y GCLM fueron considerados aprobados, con dirección favorable y techo initial_guide_candidate por ambos roles. La diferencia se limita a si ciertos resultados heterogéneos, nulos o indirectos deben clasificarse como conflicto material. FADS1 fue adjudicado; GCLM habría requerido un segundo arbitraje, lo que activó el stop estructural."
    )
    add_bullet(doc, "Calibration: 6/6 grupos con curator y critic completos; acuerdo 4/6.")
    add_bullet(doc, "Arbitraje real: uno, para FADS1; el árbitro determinó que no existía conflicto material.")
    add_bullet(doc, "Holdout: no iniciado, sin lock y sin exposición de sus seis grupos.")
    add_bullet(doc, "Costo: USD 3.071168; USD 0 en costos no observables.")
    add_bullet(doc, "Registro activo, Gold y evidence packets: sin modificaciones durante la campaña.")

    add_heading(doc, "Qué funcionó", 1)
    add_bullet(doc, "El schema semántico evitó que Sol tuviera que copiar cálculos derivados.")
    add_bullet(doc, "No reaparecieron discrepancias aritméticas ni arbitrajes disparados por scores determinísticos.")
    add_bullet(doc, "Los seis pares coincidieron en core_status, identidad, relación con el módulo, dirección dominante y techo de inferencia.")
    add_bullet(doc, "No hubo fallos HTTP, truncamientos, timeouts, retries ni costos sin telemetría.")
    add_bullet(doc, "El budget gate se mantuvo muy por debajo de USD 10 para calibration y USD 15 para campaña.")

    add_heading(doc, "Por qué se detuvo", 1)
    add_table(
        doc,
        ["Grupo", "Acuerdo", "Diferencia", "Resultado"],
        [
            ["FADS1:T1.3", "No", "Sólo material_conflict: curator Sí, critic No", "Un árbitro: No conflicto material"],
            ["GCLM:T1.6", "No", "Sólo material_conflict: curator Sí, critic No", "Segundo desacuerdo; stop antes del árbitro"],
        ],
        [1.15, 0.75, 2.8, 2.1],
        8.5,
    )
    doc.add_paragraph(
        "El protocolo permitía como máximo un desacuerdo base efectivo, porque un segundo desacuerdo vuelve imposible alcanzar 5/6. Por eso no era válido congelar el prompt ni abrir el holdout."
    )

    add_heading(doc, "Resultados por grupo", 1)
    rows = []
    for row in data["workbook"]["CALIBRATION"]:
        rows.append([row["grupo"], row["curator"], row["critic"], row["acuerdo_base"], row["arbiter"], row["decision_final"], row["resultado"]])
    add_table(doc, ["Grupo", "Curator", "Critic", "Acuerdo", "Arbiter", "Final", "Estado"], rows, [1.05, 0.9, 0.9, 0.65, 0.9, 0.9, 1.2], 7.8)

    add_heading(doc, "Hallazgo de implementación", 1)
    doc.add_paragraph(
        "La revisión post-run encontró un gap adicional: la validación inmediata de cada rol comprobó la evidencia contra la ventana seleccionada del packet, pero no contra valid_evidence_ids del Gold firmado. En FADS1, el curator utilizó siete IDs que no pertenecían a la allowlist firmada; el árbitro y la decisión final sí quedaron dentro del Gold. El gate final no lo expuso como error porque evaluó la decisión adjudicada, no cada assessment intermedio."
    )
    doc.add_paragraph(
        "Antes de una nueva calibration conviene hacer obligatorio el control por rol contra la allowlist firmada y definir con mayor precisión cuándo un resultado nulo o indirecto constituye conflicto material del mecanismo central. Esto es una corrección del proceso de evaluación, no una reapertura automática de la firma científica."
    )

    add_heading(doc, "Análisis de FADS1 y GCLM", 1)
    add_heading(doc, "FADS1:T1.3", 2)
    doc.add_paragraph(
        "Curator y critic aprobaron el grupo, asignaron initial_guide_candidate y dirección supports_relation. El curator trató evidencia mixta y resultados nulos o contextuales como conflicto material. El critic —y luego el árbitro— interpretó que esos resultados no contradicen la relación central entre FADS1 y el manejo de ácidos grasos poliinsaturados."
    )
    add_heading(doc, "GCLM:T1.6", 2)
    doc.add_paragraph(
        "Ambos roles aprobaron el grupo, mantuvieron la misma dirección y techo de inferencia. El curator consideró que resultados heterogéneos justificaban material_conflict; el critic los trató como limitaciones o evidencia no directamente opuesta al mecanismo. No se ejecutó un segundo arbitraje por el stop estructural."
    )

    add_heading(doc, "Telemetría y costos", 1)
    telemetry = data["workbook"]["TELEMETRIA_COSTOS"]
    totals = defaultdict(float)
    token_totals = Counter()
    for row in telemetry:
        totals[row["rol"]] += float(row["costo_usd"] or 0)
        for key in ("input_tokens", "cached_input_tokens", "output_tokens", "reasoning_tokens", "total_tokens"):
            token_totals[key] += int(row[key] or 0)
    add_table(
        doc,
        ["Concepto", "Valor"],
        [
            ["Llamadas científicas", len(telemetry)],
            ["Input tokens", token_totals["input_tokens"]],
            ["Cached input tokens", token_totals["cached_input_tokens"]],
            ["Output tokens", token_totals["output_tokens"]],
            ["Reasoning tokens (subconjunto del output)", token_totals["reasoning_tokens"]],
            ["Costo curator", f"USD {totals['curator']:.6f}"],
            ["Costo critic", f"USD {totals['critic']:.6f}"],
            ["Costo arbiter", f"USD {totals['arbiter']:.6f}"],
            ["Probe", f"USD {data['budget']['phase_costs_usd']['probe']:.6f}"],
            ["Total", f"USD {data['budget']['accounted_campaign_cost_usd']:.6f}"],
        ],
        [3.65, 2.3],
        9,
    )

    add_heading(doc, "Estado del holdout y controles de seguridad", 1)
    add_bullet(doc, "Holdout iniciado: no.")
    add_bullet(doc, "Prompt congelado: no.")
    add_bullet(doc, "Optimizer: deshabilitado; no se generó otra candidate.")
    add_bullet(doc, "Secreto en artefactos: el runner informó secret_scan_passed=true; el paquete técnico se vuelve a escanear antes de comprimir.")
    add_bullet(doc, f"Registro activo: {sha256(ACTIVE_REGISTRY)}.")
    add_bullet(doc, "No se ejecutaron los 105 grupos, Luna, VCF, publicación ni activación productiva.")

    add_heading(doc, "Recomendación para la siguiente autorización", 1)
    doc.add_paragraph(
        "No corresponde reutilizar ni optimizar con el holdout: sigue virgen. La siguiente campaña debería comenzar sólo después de dos ajustes verificables: (1) enforcement de valid_evidence_ids del Gold en curator, critic y arbiter antes de normalizar; y (2) una regla general que separe conflicto directo del mecanismo de heterogeneidad contextual o endpoints downstream. Luego se repiten tests y una nueva calibration independiente. Sólo si alcanza 5/6 se congela y se abre el holdout una vez."
    )

    add_heading(doc, "Cómo auditar este paquete", 1)
    add_bullet(doc, "Abrir primero el XLSX y leer RESUMEN, CALIBRATION y ARBITRAJES.")
    add_bullet(doc, "En GOLD_VS_SOL, filtrar CUMPLE=No para ver diferencias campo por campo.")
    add_bullet(doc, "En FUENTES, filtrar EN_ALLOWLIST_GOLD=No; esto muestra el gap por rol detectado en FADS1.")
    add_bullet(doc, "En TELEMETRIA_COSTOS se conservan response IDs, tokens, latencia y costo por llamada.")
    add_bullet(doc, "La auditoría técnica contiene JSON y respuestas crudas; no es necesaria para una revisión científica inicial.")

    doc.add_page_break()
    add_heading(doc, "Anexo: hashes principales", 1)
    add_table(doc, ["Artefacto", "SHA-256"], [[row["artefacto"], row["sha256"]] for row in data["workbook"]["HASHES"]], [2.15, 4.65], 7.7)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def copy_technical(data: dict[str, Any], destination: Path) -> None:
    if destination.exists():
        raise RuntimeError(f"Technical staging already exists: {destination}")
    destination.mkdir(parents=True)
    shutil.copytree(CAMPAIGN, destination / "campaign")
    inputs = destination / "inputs"
    inputs.mkdir()
    shutil.copy2(GOLD_PATH, inputs / "gold_frozen.json")
    shutil.copy2(EVIDENCE_MANIFEST_PATH, inputs / "evidence_manifest.json")
    packets_dest = inputs / "packets"
    packets_dest.mkdir()
    for card in data["gold"]["cards"]:
        shutil.copy2(packet_path(card["group_id"]), packets_dest / packet_path(card["group_id"]).name)
    implementation = destination / "implementation"
    implementation.mkdir()
    for filename in (
        "prompt_mechanism_curator_semantic_v1.md",
        "prompt_mechanism_critic_semantic_v1.md",
        "prompt_mechanism_arbiter_semantic_v1.md",
        "mechanism_curation_semantic_v1.schema.json",
        "mechanism_curation_v2.schema.json",
        "curation_v2.py",
        "run_semantic_verification_v1.py",
    ):
        shutil.copy2(SERVICE / filename, implementation / filename)


def secret_scan(root: Path) -> dict[str, Any]:
    secret = os.environ.get("HEAL_OPENAI_API_KEY", "")
    exact_hits: list[str] = []
    suspicious_hits: list[str] = []
    binary_suffixes = {".xlsx", ".docx", ".pdf", ".zip", ".png"}
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.lower() in binary_suffixes:
            continue
        raw = path.read_bytes()
        if secret and secret.encode("utf-8") in raw:
            exact_hits.append(str(path))
        text = raw.decode("utf-8", errors="ignore")
        if "Authorization: Bearer" in text or '"api_key"' in text.lower():
            suspicious_hits.append(str(path))
    return {"exact_secret_hits": exact_hits, "suspicious_header_or_field_hits": suspicious_hits, "passed": not exact_hits and not suspicious_hits}


def build_manifest(root: Path, output_path: Path) -> None:
    rows = []
    for path in sorted(item for item in root.rglob("*") if item.is_file() and item != output_path):
        rows.append({"path": path.relative_to(root).as_posix(), "bytes": path.stat().st_size, "sha256": sha256(path)})
    dump_json(output_path, {"schema_version": "semantic_v1_audit_manifest_v1", "created_at": datetime.now(timezone.utc).isoformat(), "files": rows})


def zip_tree(source: Path, destination: Path, exclude_suffixes: set[str] | None = None) -> None:
    exclude_suffixes = exclude_suffixes or set()
    with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in sorted(item for item in source.rglob("*") if item.is_file()):
            if path == destination or path.suffix.lower() in exclude_suffixes:
                continue
            archive.write(path, path.relative_to(source.parent).as_posix())


def prepare() -> None:
    if OUTPUT.exists():
        raise RuntimeError(f"Output directory already exists: {OUTPUT}")
    observed_registry = sha256(ACTIVE_REGISTRY)
    if observed_registry != EXPECTED_REGISTRY_SHA256:
        raise RuntimeError(f"Active registry hash mismatch: {observed_registry}")
    data = build_data()
    OUTPUT.mkdir(parents=True)
    human = OUTPUT / "Entregables_humanos"
    csv_dir = human / "CSVs"
    technical = OUTPUT / "Auditoria_tecnica"
    qa = OUTPUT / "_qa"
    human.mkdir()
    csv_dir.mkdir()
    qa.mkdir()

    for sheet_name, rows in data["workbook"].items():
        write_csv(csv_dir / f"{sheet_name}.csv", rows)
    dump_json(qa / "workbook_data.json", data["workbook"])
    build_docx(data, human / "Informe_SemanticV1_Calibration_2026-08-22.docx")

    guide = """# Guía práctica de lectura\n\n## Resultado en una frase\n\nLa calibration se detuvo correctamente con acuerdo 4/6; no se congeló el prompt y el holdout sigue completamente virgen.\n\n## Orden recomendado\n\n1. Abra `Informe_SemanticV1_Calibration_2026-08-22.pdf` para la explicación general.\n2. Abra el Excel y lea `RESUMEN`, `CALIBRATION` y `ARBITRAJES`.\n3. En `GOLD_VS_SOL`, filtre `CUMPLE = No` para localizar diferencias puntuales.\n4. En `FUENTES`, filtre `EN_ALLOWLIST_GOLD = No`. Es el control más importante antes de otra corrida.\n5. Use `TELEMETRIA_COSTOS` para verificar tokens, latencia, response ID y costo por llamada.\n\n## Qué revisar con un genetista\n\n- Si los resultados nulos o heterogéneos de FADS1 y GCLM contradicen directamente el mecanismo central o sólo limitan su generalización.\n- Si la evidencia marcada como conflicto cambia materialmente la interpretación individual o debe quedar como limitación contextual.\n- No hace falta revisar aritmética, tokens, hashes ni estructura JSON para esta decisión científica.\n\n## Qué no ocurrió\n\n- No se ejecutó el holdout.\n- No se congeló ningún prompt.\n- No se ejecutaron los 105 grupos, Luna ni VCF.\n- No se modificó ni publicó el registro activo.\n\n## Paquetes\n\n- `Paquete_humano_SemanticV1_2026-08-22.zip`: Excel, informe, CSVs y esta guía; no contiene JSON.\n- `Paquete_tecnico_SemanticV1_2026-08-22.zip`: respuestas crudas, decisiones, prompts, schemas, Gold, packets, telemetría y hashes. Requiere perfil técnico.\n"""
    (human / "LEEME_Guia_Practica.md").write_text(guide, encoding="utf-8")
    copy_technical(data, technical)
    dump_json(qa / "package_summary.json", {"status": data["terminal"]["status"], "holdout_started": False, "registry_sha256": observed_registry, "campaign_cost_usd": data["budget"]["accounted_campaign_cost_usd"]})
    print(json.dumps({"status": "prepared", "output": str(OUTPUT), "docx": str(human / 'Informe_SemanticV1_Calibration_2026-08-22.docx'), "workbook_data": str(qa / 'workbook_data.json')}, ensure_ascii=False))


def finalize() -> None:
    human = OUTPUT / "Entregables_humanos"
    technical = OUTPUT / "Auditoria_tecnica"
    required = [
        human / "LLM1_Tier1_Sol_SemanticV1_2026-08-22.xlsx",
        human / "Informe_SemanticV1_Calibration_2026-08-22.docx",
        human / "Informe_SemanticV1_Calibration_2026-08-22.pdf",
        human / "LEEME_Guia_Practica.md",
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise RuntimeError("Missing required deliverables: " + ", ".join(missing))
    scan = secret_scan(technical)
    dump_json(OUTPUT / "_qa" / "secret_scan.json", scan)
    if not scan["passed"]:
        raise RuntimeError("Secret scan failed; see _qa/secret_scan.json")
    build_manifest(technical, technical / "manifest_sha256.json")
    technical_zip = OUTPUT / "Paquete_tecnico_SemanticV1_2026-08-22.zip"
    human_zip = OUTPUT / "Paquete_humano_SemanticV1_2026-08-22.zip"
    zip_tree(technical, technical_zip)
    with zipfile.ZipFile(human_zip, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in sorted(item for item in human.rglob("*") if item.is_file()):
            if path.suffix.lower() == ".json":
                continue
            archive.write(path, path.relative_to(human.parent).as_posix())
    build_manifest(OUTPUT, OUTPUT / "manifest_entrega_sha256.json")
    print(json.dumps({"status": "finalized", "human_zip": str(human_zip), "technical_zip": str(technical_zip), "secret_scan_passed": scan["passed"]}, ensure_ascii=False))


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("action", choices=["prepare", "finalize"])
    args = parser.parse_args()
    if args.action == "prepare":
        prepare()
    else:
        finalize()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
