#!/usr/bin/env python3
"""Build the human-readable Tier 1 reconciliation report."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5B6673"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GREEN = "E7F4EC"
PALE_YELLOW = "FFF6E0"
WHITE = "FFFFFF"
PAGE_WIDTH_DXA = 9360


def set_font(run, *, size=11, color="000000", bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def shade(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        properties.append(element)
    element.set(qn("w:fill"), color)


def cell_margins(cell, top=80, bottom=80, start=120, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_geometry(table, widths, indent=120):
    table.autofit = False
    properties = table._tbl.tblPr
    width = properties.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    table_indent = properties.first_child_found_in("w:tblInd")
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        properties.append(table_indent)
    table_indent.set(qn("w:w"), str(indent))
    table_indent.set(qn("w:type"), "dxa")
    for index, item in enumerate(widths):
        table._tbl.tblGrid.gridCol_lst[index].set(qn("w:w"), str(item))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell._tc.tcPr.tcW.set(qn("w:w"), str(widths[index]))
            cell._tc.tcPr.tcW.set(qn("w:type"), "dxa")
            cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def paragraph(doc, text="", *, size=11, color="000000", bold=False, italic=False, before=0, after=6, line=1.10, align=None):
    item = doc.add_paragraph()
    item.paragraph_format.space_before = Pt(before)
    item.paragraph_format.space_after = Pt(after)
    item.paragraph_format.line_spacing = line
    if align is not None:
        item.alignment = align
    set_font(item.add_run(text), size=size, color=color, bold=bold, italic=italic)
    return item


def heading(doc, text, level=1):
    item = doc.add_paragraph(style=f"Heading {level}")
    item.paragraph_format.space_before = Pt({1: 16, 2: 12, 3: 8}[level])
    item.paragraph_format.space_after = Pt({1: 8, 2: 6, 3: 4}[level])
    item.paragraph_format.keep_with_next = True
    run = item.add_run(text)
    set_font(run, size={1: 16, 2: 13, 3: 12}[level], color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level], bold=True)
    return item


def bullet(doc, text):
    item = doc.add_paragraph(style="List Bullet")
    item.paragraph_format.left_indent = Inches(0.5)
    item.paragraph_format.first_line_indent = Inches(-0.25)
    item.paragraph_format.space_after = Pt(8)
    item.paragraph_format.line_spacing = 1.167
    set_font(item.add_run(text), size=11)


def numbered(doc, text):
    item = doc.add_paragraph(style="List Number")
    item.paragraph_format.left_indent = Inches(0.5)
    item.paragraph_format.first_line_indent = Inches(-0.25)
    item.paragraph_format.space_after = Pt(8)
    item.paragraph_format.line_spacing = 1.167
    set_font(item.add_run(text), size=11)


def callout(doc, label, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table_geometry(table, [PAGE_WIDTH_DXA])
    cell = table.cell(0, 0)
    shade(cell, fill)
    body = cell.paragraphs[0]
    body.paragraph_format.space_after = Pt(0)
    body.paragraph_format.line_spacing = 1.10
    set_font(body.add_run(label + " "), color=NAVY, bold=True)
    set_font(body.add_run(text))
    paragraph(doc, "", after=2)


def add_table(doc, headers, rows, widths, *, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade(cell, LIGHT_BLUE)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        set_font(cell.paragraphs[0].add_run(header), size=9.5, color=NAVY, bold=True)
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].paragraphs[0].paragraph_format.space_after = Pt(0)
            cells[index].paragraphs[0].paragraph_format.line_spacing = 1.05
            set_font(cells[index].paragraphs[0].add_run(str(value)), size=font_size)
    table_geometry(table, widths)
    paragraph(doc, "", after=3)
    return table


def footer(section):
    item = section.footer.paragraphs[0]
    item.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    item.paragraph_format.space_after = Pt(0)
    set_font(item.add_run("HEAL by FON | Recuración Tier 1 | Uso interno"), size=8.5, color=MUTED)


def build(foundation, old_foundation, manifest):
    groups = [foundation["groups"][group] for group in foundation["groups"]]
    packet_by_group = {
        row["group_id"]: json.loads(Path(row["packet_path"]).read_text(encoding="utf-8"))
        for row in manifest["packets"]
    }
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    footer(section)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    paragraph(doc, "INFORME DE DECISIÓN", size=10, color=MUTED, bold=True, after=3)
    title = paragraph(doc, "Reconciliación Tier 1 lista para firma humana", size=23, color="000000", bold=True, after=4)
    title.paragraph_format.keep_with_next = True
    paragraph(doc, "Qué se reparó, qué quedó validado y cómo cerrar el gold de 12 grupos", size=14, color=MUTED, after=14)
    for label, value in [
        ("Fecha", "14/08/2026"), ("Estado", "12/12 técnicamente reconciliados; firma pendiente"),
        ("Cutoff de publicaciones", "28/07/2026"), ("Alcance", "Curación científica interna; sin activación de Luna"),
    ]:
        item = doc.add_paragraph()
        item.paragraph_format.space_after = Pt(2)
        set_font(item.add_run(label + ": "), bold=True)
        set_font(item.add_run(value))
    paragraph(doc, "", after=6)
    callout(doc, "Resultado principal.", "Los 12 paquetes superan los controles técnicos. El único bloqueo restante es deliberado: falta tu aprobación explícita de las decisiones científicas.", PALE_GREEN)
    callout(doc, "Qué no ocurrió.", "No se ejecutó GPT-5.6 Sol, no se ejecutó Luna, no se cambió el registro activo y no se regeneró ningún VCF.", PALE_YELLOW)

    heading(doc, "1. Resultado ejecutivo")
    metrics = [
        ("12/12", "paquetes reconciliados"), ("0", "bloqueos técnicos"),
        ("12", "firmas pendientes"), ("20", "fuentes seleccionadas por grupo"),
    ]
    add_table(doc, ["Resultado", "Significado"], metrics, [1800, 7560], font_size=10.5)
    paragraph(doc, "Esto significa que la base ya es consistente para una revisión humana final. No significa todavía que el gold esté aprobado ni que el sistema pueda comenzar la calibración automática.")

    heading(doc, "2. Qué se corrigió")
    for text in [
        "Los PMID que la revisión había marcado como válidos ahora se recuperan de forma explícita, aunque no aparezcan entre los primeros resultados de una búsqueda general.",
        "Las fuentes marcadas como inválidas se conservan para auditoría, pero quedan excluidas y no pueden entrar en la ventana que ve el modelo.",
        "PubMed pasó a ser la fuente prioritaria para título y DOI cuando existe un PMID. Europe PMC sólo completa registros ausentes, evitando falsos conflictos de metadatos.",
        "La selección respeta el máximo de 20 fuentes y no permite que una fuente obligatoria saltee el cutoff, sea una review usada como prueba primaria o duplique una publicación de la misma línea.",
        "La identidad de PEMT quedó corregida a NCBI Gene 10400. El identificador 4582 corresponde a MUC1 y ya no puede sostener la identidad de PEMT.",
        "Los problemas técnicos quedaron separados del estado científico. Una falla de fuente no baja automáticamente un grupo a withheld ni se interpreta como evidencia negativa.",
    ]:
        bullet(doc, text)

    heading(doc, "3. Estado de los 12 grupos")
    rows = []
    for group in groups:
        packet = packet_by_group[group["group_id"]]
        rows.append((
            group["group_id"], "Calibración" if group["split"] == "calibration" else "Holdout",
            group["scientific_status_proposed"], group["inference_ceiling_proposed"],
            len(group["allowed_evidence_ids_proposed"]), packet["selection_summary"]["selected_total"],
        ))
    add_table(doc, ["Grupo", "Uso", "Estado propuesto", "Techo", "Fuentes válidas", "En ventana"], rows, [1350, 1250, 1900, 2380, 1350, 1130], font_size=8.7)
    paragraph(doc, "Los estados siguen siendo propuestas. La planilla de firma permite aprobarlos o corregirlos sin alterar los archivos técnicos originales.", size=9.5, color=MUTED, italic=True)

    doc.add_page_break()
    heading(doc, "4. Ejemplos para orientar la revisión")
    examples = [
        ("PEMT:T1.3", "Identidad", "Confirmar que la base usa NCBI Gene 10400 y que 4582/MUC1 aparece sólo como identidad descartada."),
        ("CYCS:T1.1", "Conflicto", "La relación con energía mitocondrial es directa; el conflicto debe conservarse entre ensayos/modelos sin convertirlo en diagnóstico."),
        ("IL6:T1.4", "Contexto inmune", "La relación es utilizable, pero señalización clásica y trans pueden diferir. No inferir inflamación crónica o riesgo individual."),
        ("ABCB1:T1.6", "Farmacogenómica", "Puede explicarse el contexto por sustrato, pero no recomendar cambios de medicación ni una capacidad global de 'detox'."),
        ("COL14A1:T1.5", "Techo prudente", "La biología de matriz/tendón es útil como contexto, pero la traducción a lesión individual no alcanza: context_only."),
    ]
    add_table(doc, ["Grupo", "Qué ilustra", "Qué corroborar"], examples, [1500, 1700, 6160], font_size=9.2)

    heading(doc, "5. Guía práctica para aprobar o corregir")
    for text in [
        "Abrí la hoja FIRMA_12_GRUPOS del Excel. Cada fila corresponde a un grupo.",
        "Revisá estado, posibilidad de usar contexto, techo de inferencia, dirección, conflictos, límites y fuentes incluidas/excluidas.",
        "Si estás de acuerdo, elegí APROBADO y completá nombre y fecha. Si no, elegí CORREGIR y describí el cambio en la última columna.",
        "No hace falta reescribir una respuesta ideal. Sólo se aprueba el comportamiento aceptable y la evidencia que puede sostenerlo.",
        "Devolvé la planilla completa. El compilador verificará nuevamente hashes, fuentes y los 12 grupos antes de crear el gold ejecutable.",
    ]:
        numbered(doc, text)
    callout(doc, "Cuándo sumar un genetista.", "Conviene pedir una segunda mirada si existe duda de identidad, si una asociación humana parece venir de otra condición o población, si dos estudios importantes chocan, o antes de habilitar initial_guide_candidate. No es necesario derivar cada fila.", LIGHT_GRAY)

    heading(doc, "6. Qué sucede después de la firma")
    for text in [
        "Se compila un gold inmutable y se valida que las 12 tarjetas estén firmadas, fechadas y vinculadas al hash correcto.",
        "Sol ejecuta dos evaluaciones independientes: primero los seis casos de calibración y luego los seis casos holdout.",
        "Si una versión no cumple 12/12, se ajusta sólo el prompt general. El prompt no puede contener nombres de los genes del gold.",
        "Recién con una versión aprobada se generan los 105 paquetes Tier 1 y se ejecutan las 210 llamadas base más adjudicaciones.",
        "El snapshot de 105 grupos vuelve a revisión humana antes de cualquier publicación o prueba con los dos VCF.",
    ]:
        bullet(doc, text)

    heading(doc, "7. Controles que siguen vigentes")
    for text in [
        "Nunca convertir draft en withheld automáticamente.",
        "Una base autoritativa sola no alcanza para aprobar.",
        "Una review sirve para descubrir evidencia primaria, no como único sostén.",
        "La ausencia de datos no significa benignidad, homocigosis de referencia ni falta de relevancia.",
        "La curación del mecanismo no depende de las variantes presentes en los dos VCF.",
        "Un estado approved no habilita diagnóstico, tratamiento, dosis, suplementación ni farmacogenómica accionable.",
    ]:
        bullet(doc, text)

    heading(doc, "8. Archivos para la revisión")
    add_table(doc, ["Archivo", "Uso"], [
        ("LLM1_Tier1_Reconciliacion_y_Firma_2026-08-14.xlsx", "Documento principal: resumen, firma, fuentes válidas/excluidas, diferencias y glosario."),
        ("LLM1_Tier1_Informe_Reconciliacion_2026-08-14.docx", "Este informe en formato editable."),
        ("LLM1_Tier1_Informe_Reconciliacion_2026-08-14.pdf", "Versión de lectura y diseminación."),
        ("human_signoff_template.csv", "Respaldo tabular de la hoja de firma; útil para importación técnica posterior."),
    ], [4300, 5060], font_size=9.5)

    heading(doc, "9. Trazabilidad técnica mínima")
    add_table(doc, ["Elemento", "Valor"], [
        ("Fundación reconciliada", foundation["snapshot_sha256"]),
        ("Manifest de evidencia", manifest["manifest_sha256"]),
        ("Fundación anterior", old_foundation["snapshot_sha256"]),
        ("Registro activo modificado", "No"),
        ("Estado siguiente", "Firma humana explícita de 12/12"),
    ], [2500, 6860], font_size=8.7)
    paragraph(doc, "El material constituye curación científica interna y una base de evaluación. No equivale a validación genética, bioinformática o clínica independiente.", size=9.5, color=MUTED, italic=True, before=8)
    return doc


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--foundation", required=True)
    parser.add_argument("--old-foundation", required=True)
    parser.add_argument("--evidence-manifest", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    foundation = json.loads(Path(args.foundation).read_text(encoding="utf-8"))
    old = json.loads(Path(args.old_foundation).read_text(encoding="utf-8"))
    manifest = json.loads(Path(args.evidence_manifest).read_text(encoding="utf-8"))
    document = build(foundation, old, manifest)
    document.core_properties.title = "Reconciliación Tier 1 lista para firma humana"
    document.core_properties.subject = "Curación científica interna de LLM1"
    document.core_properties.author = "HEAL by FON"
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(json.dumps({"output": str(output.resolve()), "status": "created"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
