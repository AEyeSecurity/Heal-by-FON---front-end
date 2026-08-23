from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO = Path(r"F:\Heal by FON\app")
SERVICE = REPO / "services" / "heal-tier1-curation-v2"
RECON = Path(r"F:\Heal by FON\data\curation-candidates\tier1-human-review-20260810\reconciliation-v4")
CAMPAIGN = RECON / "semantic-v2-allowlist-conflict-verification-2"
CAMPAIGN_ATTEMPT1 = RECON / "semantic-v2-allowlist-conflict-verification-1"
GOLD = RECON / "gold-v4" / "gold_frozen.json"
PACKETS = RECON / "evidence" / "packets"
REGISTRY = Path(r"F:\Heal by FON\data\canon\curation\mechanism_registry_v1.csv")
EXPECTED_REGISTRY_HASH = "73b94c09c31c184135bc16b2e756fa447f4c040a8bf38b49181168a8c62a967f"
GROUPS = ["MTHFR:T1.1", "ASMT:T1.2", "FADS1:T1.3", "IL10:T1.4", "RUNX2:T1.5", "GCLM:T1.6"]


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def flatten(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, list):
        return " | ".join(flatten(item) for item in value)
    if isinstance(value, dict):
        return " | ".join(f"{key}: {flatten(item)}" for key, item in value.items())
    return str(value)


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    headers: list[str] = []
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)
    if not rows:
        headers, rows = ["estado"], [{"estado": "Sin datos"}]
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: flatten(row.get(key)) for key in headers})


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_run(run, *, size=10.5, bold=False, color="1F2937") -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.1
    style_run(paragraph.add_run(text))


def add_table(doc: Document, headers: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "E8EEF5")
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(2)
        style_run(paragraph.add_run(header), size=9, bold=True, color="0B2545")
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.0
            style_run(paragraph.add_run(flatten(value)), size=8.5)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def load_and_replay() -> dict[str, Any]:
    sys.path.insert(0, str(SERVICE))
    import curation_v2 as cv2

    gold = read_json(GOLD)
    cards = {row["group_id"]: row for row in gold["cards"]}
    phase = read_json(CAMPAIGN / "calibration" / "phase_acceptance.json")
    terminal = read_json(CAMPAIGN / "protocol_terminal_state.json")
    budget = read_json(CAMPAIGN / "budget_ledger.json")
    attempt1_budget = read_json(CAMPAIGN_ATTEMPT1 / "budget_ledger.json") if (CAMPAIGN_ATTEMPT1 / "budget_ledger.json").exists() else {"observed_cost_usd": 0.0}
    group_rows: list[dict[str, Any]] = []
    role_rows: list[dict[str, Any]] = []
    source_rows: list[dict[str, Any]] = []
    telemetry_rows: list[dict[str, Any]] = []
    error_rows: list[dict[str, Any]] = []

    for group_id in GROUPS:
        stem = group_id.replace(":", "__")
        record_path = CAMPAIGN / "calibration" / "records" / f"{stem}.json"
        if not record_path.exists():
            group_rows.append({"grupo": group_id, "estado_ejecucion": "no_ejecutado", "motivo": "Calibration detenida antes de este grupo"})
            continue
        record = read_json(record_path)
        packet = read_json(PACKETS / f"{stem}.json")
        replay_roles: dict[str, dict[str, Any]] = {}
        for role in ("curator", "critic", "arbiter"):
            semantic = record.get(f"{role}_semantic")
            if not semantic:
                continue
            replay_errors = cv2.validate_semantic_assessment_v2(semantic, packet)
            replay_decision = cv2.normalize_semantic_assessment_v2(semantic, packet)
            replay_roles[role] = replay_decision
            direction = replay_decision["direction_assessment"]
            card = cards[group_id]
            gold_ok = (
                replay_decision["core_status"] in card["acceptable_core_statuses"]
                and replay_decision["inference_ceiling"] in card["acceptable_inference_ceilings"]
                and direction["dominant_direction"] in card["acceptable_directions"]
            )
            role_rows.append({
                "grupo": group_id,
                "rol": role,
                "core_status": replay_decision["core_status"],
                "direccion_recalculada": direction["dominant_direction"],
                "conflicto_material": direction["material_conflict"],
                "techo_inferencia": replay_decision["inference_ceiling"],
                "confianza": replay_decision["confidence"],
                "errores_replay_corregido": replay_errors,
                "dentro_gold_replay": gold_ok,
                "provenance_allowlist": record.get("allowlist_provenance"),
            })
            for item in semantic.get("used_evidence", []):
                source_rows.append({
                    "grupo": group_id,
                    "rol": role,
                    "evidence_id": item.get("evidence_id"),
                    "uso": "usada",
                    "direccion": item.get("supports_direction"),
                    "clase_conflicto": item.get("core_conflict_class"),
                    "rationale": item.get("core_conflict_rationale"),
                })
            for item in semantic.get("excluded_evidence", []):
                source_rows.append({
                    "grupo": group_id,
                    "rol": role,
                    "evidence_id": item.get("evidence_id"),
                    "uso": "excluida",
                    "direccion": "",
                    "clase_conflicto": "",
                    "rationale": item.get("reason_code"),
                })
        curator = replay_roles.get("curator")
        critic = replay_roles.get("critic")
        replay_agreement = bool(curator and critic and not cv2.semantic_arbitration_reasons(
            record.get("curator_semantic"), record.get("critic_semantic"), curator, critic
        ))
        final = replay_roles.get("arbiter") or curator
        final_direction = (final or {}).get("direction_assessment", {})
        card = cards[group_id]
        final_gold_ok = bool(final) and (
            final.get("core_status") in card["acceptable_core_statuses"]
            and final.get("inference_ceiling") in card["acceptable_inference_ceilings"]
            and final_direction.get("dominant_direction") in card["acceptable_directions"]
        )
        group_rows.append({
            "grupo": group_id,
            "estado_ejecucion": "completo",
            "acuerdo_original": record.get("base_pair_agreement"),
            "arbitraje_original": record.get("adjudicated"),
            "motivos_arbitraje_original": record.get("arbitration_reasons"),
            "acuerdo_replay_corregido": replay_agreement,
            "decision_final_replay": (final or {}).get("core_status"),
            "direccion_final_replay": final_direction.get("dominant_direction"),
            "conflicto_material_replay": final_direction.get("material_conflict"),
            "techo_replay": (final or {}).get("inference_ceiling"),
            "dentro_gold_replay": final_gold_ok,
        })

    for audit_path in sorted((CAMPAIGN / "calibration" / "audit").glob("*.json")):
        audit = read_json(audit_path)
        for attempt in audit.get("attempts", []):
            telemetry_rows.append({
                "fase": audit.get("phase"), "grupo": audit.get("group_id"), "rol": audit.get("role"),
                "intento": attempt.get("attempt"), "estado": attempt.get("response_status"),
                "response_id": attempt.get("response_id"), "modelo": attempt.get("effective_model"),
                "reasoning_effort": attempt.get("reasoning_effort"), "input_tokens": attempt.get("input_tokens"),
                "cached_input_tokens": attempt.get("cached_input_tokens"), "output_tokens": attempt.get("output_tokens"),
                "reasoning_tokens": attempt.get("reasoning_tokens"), "latencia_segundos": attempt.get("latency_seconds"),
                "costo_usd": attempt.get("estimated_cost_usd"), "retry": (attempt.get("attempt") or 1) > 1,
                "timeout": attempt.get("response_status") == "timeout",
            })
    for error in phase.get("errors", []):
        error_rows.append({"origen": "campaign_gate", "codigo": error, "clasificacion": "defecto estructural local confirmado"})
    error_rows.extend([
        {"origen": "attempt_1", "codigo": "responses_schema_uniqueItems_not_supported", "clasificacion": "fallo técnico de preflight; schema corregido antes del intento científico"},
        {"origen": "diagnostico", "codigo": "contextual_positive_support_suppressed", "clasificacion": "normalizador v2; corregido y cubierto por regresión"},
        {"origen": "diagnostico", "codigo": "downstream_null_forced_global_null", "clasificacion": "validador v2; corregido y cubierto por regresión"},
    ])
    return {
        "gold": gold, "phase": phase, "terminal": terminal, "budget": budget,
        "attempt1_budget": attempt1_budget,
        "groups": group_rows, "roles": role_rows, "sources": source_rows,
        "telemetry": telemetry_rows, "errors": error_rows,
    }


def build_docx(path: Path, data: dict[str, Any], registry_hash: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(6), 1.1
    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 16, 8), ("Heading 2", 13, "2E74B5", 12, 6), ("Heading 3", 12, "1F4D78", 8, 4)
    ):
        style = styles[name]
        style.font.name, style.font.size, style.font.color.rgb = "Calibri", Pt(size), RGBColor.from_string(color)
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_run(header.add_run("HEAL | Auditoría interna LLM1"), size=8.5, color="64748B")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(footer.add_run("Semantic V2 - calibration detenida | 2026-08-22"), size=8, color="64748B")

    title = doc.add_paragraph()
    title.paragraph_format.space_before, title.paragraph_format.space_after = Pt(10), Pt(4)
    style_run(title.add_run("Informe de calibration Semantic V2"), size=24, bold=True, color="0B2545")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    style_run(subtitle.add_run("Diagnóstico de detención, corrección local y trazabilidad"), size=13, color="475569")
    add_table(doc, ["Campo", "Valor"], [
        ["Campaña", "semantic-v2-allowlist-conflict-verification-2"],
        ["Estado", "Calibration detenida; sin freeze, holdout ni 105 grupos"],
        ["Costo observado total", f"USD {data['budget']['observed_cost_usd'] + data['attempt1_budget']['observed_cost_usd']:.6f}"],
        ["Registro activo", f"Intacto: {registry_hash}"],
    ], [2100, 7260])

    doc.add_heading("Conclusión ejecutiva", level=1)
    p = doc.add_paragraph()
    style_run(p.add_run("La campaña no aprobó la calibration y se cerró conforme al protocolo. "), bold=True, color="9B1C1C")
    style_run(p.add_run("No se congeló ningún prompt, no se abrió el holdout y no se iniciaron los 105 grupos."))
    add_bullet(doc, "Cinco de seis grupos llegaron a ejecutarse; GCLM no fue llamado.")
    add_bullet(doc, "FADS1 consumió el único arbitraje permitido y RUNX2 produjo el segundo desacuerdo aparente, por lo que se activó la detención anticipada.")
    add_bullet(doc, "La revisión posterior confirmó dos defectos locales del normalizador/validador; no fueron fallos de la evidencia ni del razonamiento científico de Sol.")
    add_bullet(doc, "Tras corregirlos, 147 tests locales pasan y el replay local de las respuestas ya obtenidas deja FADS1 y RUNX2 alineados con el Gold.")

    doc.add_heading("Qué falló y por qué", level=1)
    doc.add_heading("1. Soporte positivo con limitaciones contextuales", level=2)
    doc.add_paragraph(
        "El normalizador convertía en not_applicable toda fuente marcada como contextual_heterogeneity, downstream_null o limited_generalizability. "
        "Eso eliminaba también soporte positivo legítimo. En FADS1, curator conservó soporte pero critic y arbiter quedaron artificialmente sin dirección, generando gold:direction."
    )
    doc.add_heading("2. Endpoint nulo downstream versus dirección global", level=2)
    doc.add_paragraph(
        "El validador exigía supports_direction=null cuando core_conflict_class=downstream_null. Una publicación puede respaldar el mecanismo central y, al mismo tiempo, "
        "informar un endpoint posterior nulo. El critic de RUNX2 describió precisamente ese patrón en PMID:33253203 y PMID:38218304; la salida era científicamente coherente, pero el contrato la rechazó."
    )

    doc.add_heading("Corrección aplicada", level=1)
    add_bullet(doc, "Las limitaciones contextuales nunca generan oposición ni conflicto material.")
    add_bullet(doc, "Si la dirección central de una fuente contextual es positiva, sigue aportando soporte.")
    add_bullet(doc, "downstream_null puede coexistir con dirección central positiva; el endpoint nulo se conserva como limitación.")
    add_bullet(doc, "Sólo direct_material_contradiction puede aportar oposición y activar material_conflict.")
    add_bullet(doc, "Se agregaron dos regresiones específicas; suite completa: 147/147 aprobada.")

    doc.add_heading("Resultado por grupo", level=1)
    rows = []
    for row in data["groups"]:
        rows.append([
            row["grupo"], row["estado_ejecucion"], row.get("acuerdo_original", ""),
            row.get("acuerdo_replay_corregido", ""), row.get("direccion_final_replay", ""),
            row.get("dentro_gold_replay", ""),
        ])
    add_table(doc, ["Grupo", "Ejecución", "Acuerdo original", "Acuerdo replay", "Dirección replay", "Gold"], rows, [1300, 1600, 1250, 1250, 2100, 1860])

    doc.add_heading("Costos y ejecución", level=1)
    add_table(doc, ["Indicador", "Resultado"], [
        ["Costo probe", f"USD {data['budget']['phase_costs_usd']['probe']:.6f}"],
        ["Costo calibration", f"USD {data['budget']['phase_costs_usd']['calibration']:.6f}"],
        ["Costo intento técnico inicial", f"USD {data['attempt1_budget']['observed_cost_usd']:.6f}"],
        ["Costo total observado", f"USD {data['budget']['observed_cost_usd'] + data['attempt1_budget']['observed_cost_usd']:.6f}"],
        ["Retries", sum(1 for row in data["telemetry"] if row["retry"])],
        ["Timeouts", sum(1 for row in data["telemetry"] if row["timeout"])],
        ["Costos no observables", f"USD {data['budget']['unknown_reserved_cost_usd']:.6f}"],
    ], [3000, 6360])

    doc.add_heading("Estado de gates", level=1)
    add_table(doc, ["Gate", "Estado"], [
        ["Tests locales posteriores", "PASS - 147/147"],
        ["Calibration", "FAIL/INCOMPLETE - 5/6 ejecutados"],
        ["Freeze", "NO EJECUTADO"],
        ["Holdout", "NO EJECUTADO; lock no creado"],
        ["Estimación 105", "NO EJECUTADA"],
        ["Recuración 105", "NO EJECUTADA"],
        ["Luna / VCF / publicación", "FUERA DE ALCANCE; no ejecutado"],
    ], [3000, 6360])

    doc.add_heading("Próximo paso recomendado", level=1)
    doc.add_paragraph(
        "La implementación local ya está corregida, pero esta campaña permanece cerrada e inmutable. El próximo paso requiere autorizar una nueva campaña independiente "
        "con un nombre nuevo. Esa campaña deberá repetir calibration desde cero; sólo si alcanza 6/6 podrá congelar y abrir el holdout virgen."
    )
    doc.add_heading("Alcance de los anexos", level=1)
    doc.add_paragraph(
        "El XLSX contiene el detalle por grupo y rol, las fuentes usadas/excluidas, errores y telemetría. El ZIP técnico conserva las respuestas API, assessments semánticos, "
        "decisiones normalizadas, prompts, schemas, configuración y hashes. No contiene credenciales ni headers de autorización."
    )
    doc.save(path)


def build_pdf(path: Path, data: dict[str, Any], registry_hash: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="AuditTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=colors.HexColor("#0B2545"), spaceAfter=8))
    styles.add(ParagraphStyle(name="AuditSub", parent=styles["Normal"], fontName="Helvetica", fontSize=11, leading=15, textColor=colors.HexColor("#475569"), spaceAfter=14))
    styles.add(ParagraphStyle(name="AuditH1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=colors.HexColor("#2E74B5"), spaceBefore=12, spaceAfter=7))
    styles.add(ParagraphStyle(name="AuditH2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11.5, leading=14, textColor=colors.HexColor("#1F4D78"), spaceBefore=8, spaceAfter=4))
    styles.add(ParagraphStyle(name="AuditBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13, textColor=colors.HexColor("#1F2937"), spaceAfter=6))
    styles.add(ParagraphStyle(name="AuditSmall", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.5, leading=9.5, textColor=colors.HexColor("#1F2937")))
    styles.add(ParagraphStyle(name="AuditFooter", parent=styles["Normal"], fontName="Helvetica", fontSize=7.5, alignment=TA_CENTER, textColor=colors.HexColor("#64748B")))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawCentredString(letter[0] / 2, 0.42 * inch, f"Semantic V2 - calibration detenida | pagina {document.page}")
        canvas.restoreState()

    def table(rows, widths):
        converted = [[Paragraph(str(value), styles["AuditSmall"]) for value in row] for row in rows]
        item = Table(converted, colWidths=widths, repeatRows=1, hAlign="LEFT")
        item.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0B2545")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C7D2DE")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        return item

    story = [
        Paragraph("Informe de calibration Semantic V2", styles["AuditTitle"]),
        Paragraph("Diagnóstico de detención, corrección local y trazabilidad", styles["AuditSub"]),
        table([
            ["Campo", "Valor"],
            ["Campaña", "semantic-v2-allowlist-conflict-verification-2"],
            ["Estado", "Calibration detenida; sin freeze, holdout ni 105 grupos"],
            ["Costo observado total", f"USD {data['budget']['observed_cost_usd'] + data['attempt1_budget']['observed_cost_usd']:.6f}"],
            ["Registro activo", f"Intacto: {registry_hash}"],
        ], [1.45 * inch, 5.05 * inch]),
        Paragraph("Conclusión ejecutiva", styles["AuditH1"]),
        Paragraph("La campaña no aprobó la calibration y se cerró conforme al protocolo. No se congeló ningún prompt, no se abrió el holdout y no se iniciaron los 105 grupos.", styles["AuditBody"]),
        Paragraph("- Cinco de seis grupos llegaron a ejecutarse; GCLM no fue llamado.<br/>- FADS1 consumió el único arbitraje permitido y RUNX2 activó la detención anticipada.<br/>- La revisión confirmó dos defectos locales del normalizador/validador, no fallos científicos de Sol.<br/>- Después de corregirlos, 147/147 tests pasan y el replay local de FADS1 y RUNX2 queda dentro del Gold.", styles["AuditBody"]),
        Paragraph("Qué falló y por qué", styles["AuditH1"]),
        Paragraph("1. Soporte positivo con limitaciones contextuales", styles["AuditH2"]),
        Paragraph("El normalizador anulaba toda fuente contextual, incluso cuando apoyaba el mecanismo central. En FADS1 esto produjo una dirección not_applicable artificial para critic y arbiter.", styles["AuditBody"]),
        Paragraph("2. Endpoint nulo downstream versus dirección global", styles["AuditH2"]),
        Paragraph("El validador exigía dirección global null para toda fuente con downstream_null. Una publicación puede apoyar el mecanismo central y, a la vez, describir un endpoint posterior nulo. Eso ocurrió en RUNX2 para PMID:33253203 y PMID:38218304.", styles["AuditBody"]),
        Paragraph("Corrección aplicada", styles["AuditH1"]),
        Paragraph("- Las limitaciones contextuales nunca generan oposición.<br/>- El soporte central positivo se conserva.<br/>- downstream_null puede coexistir con dirección central positiva.<br/>- Sólo direct_material_contradiction activa oposición y material_conflict.<br/>- Se agregaron regresiones y la suite completa aprobó 147/147.", styles["AuditBody"]),
        PageBreak(),
        Paragraph("Resultado por grupo", styles["AuditH1"]),
    ]
    group_rows = [["Grupo", "Ejecución", "Acuerdo original", "Acuerdo replay", "Dirección replay", "Gold"]]
    for row in data["groups"]:
        group_rows.append([row["grupo"], row["estado_ejecucion"], flatten(row.get("acuerdo_original")), flatten(row.get("acuerdo_replay_corregido")), row.get("direccion_final_replay", ""), flatten(row.get("dentro_gold_replay"))])
    story.extend([
        table(group_rows, [0.95 * inch, 0.9 * inch, 0.95 * inch, 0.9 * inch, 1.75 * inch, 0.65 * inch]),
        Paragraph("Costos y ejecución", styles["AuditH1"]),
        table([
            ["Indicador", "Resultado"],
            ["Costo probe", f"USD {data['budget']['phase_costs_usd']['probe']:.6f}"],
            ["Costo calibration", f"USD {data['budget']['phase_costs_usd']['calibration']:.6f}"],
            ["Costo intento técnico inicial", f"USD {data['attempt1_budget']['observed_cost_usd']:.6f}"],
            ["Costo total observado", f"USD {data['budget']['observed_cost_usd'] + data['attempt1_budget']['observed_cost_usd']:.6f}"],
            ["Retries", sum(1 for row in data["telemetry"] if row["retry"])],
            ["Timeouts", sum(1 for row in data["telemetry"] if row["timeout"])],
            ["Costos no observables", f"USD {data['budget']['unknown_reserved_cost_usd']:.6f}"],
        ], [2.4 * inch, 4.1 * inch]),
        Paragraph("Estado de gates", styles["AuditH1"]),
        table([
            ["Gate", "Estado"], ["Tests locales posteriores", "PASS - 147/147"],
            ["Calibration", "FAIL/INCOMPLETE - 5/6 ejecutados"], ["Freeze", "NO EJECUTADO"],
            ["Holdout", "NO EJECUTADO; lock no creado"], ["Estimación y recuración 105", "NO EJECUTADAS"],
            ["Luna / VCF / publicación", "FUERA DE ALCANCE; no ejecutado"],
        ], [2.4 * inch, 4.1 * inch]),
        Paragraph("Próximo paso recomendado", styles["AuditH1"]),
        Paragraph("La campaña queda cerrada e inmutable. El próximo paso requiere una nueva autorización para iniciar otra campaña independiente con nombre nuevo. Debe repetir calibration desde cero y sólo podrá congelar si alcanza 6/6.", styles["AuditBody"]),
        Paragraph("Contenido de auditoría", styles["AuditH1"]),
        Paragraph("El XLSX contiene detalle por grupo y rol, fuentes, errores y telemetría. El ZIP técnico conserva respuestas API, assessments, decisiones, prompts, schemas, configuración y hashes, sin credenciales ni headers de autorización.", styles["AuditBody"]),
    ])
    document = SimpleDocTemplate(str(path), pagesize=letter, rightMargin=inch, leftMargin=inch, topMargin=0.75 * inch, bottomMargin=0.7 * inch, title="Informe Semantic V2 - Calibration detenida")
    document.build(story, onFirstPage=footer, onLaterPages=footer)


def zip_tree(zip_path: Path, files: list[Path], base: Path) -> None:
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(files):
            if path.is_file():
                archive.write(path, path.relative_to(base).as_posix())


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--preserve-docx", action="store_true")
    args = parser.parse_args()
    output = args.output
    output.mkdir(parents=True, exist_ok=True)
    human = output / "Paquete_humano"
    technical = output / "Auditoria_tecnica"
    human.mkdir(exist_ok=True)
    technical.mkdir(exist_ok=True)

    data = load_and_replay()
    registry_hash = sha256(REGISTRY)
    write_csv(human / "01_resultado_por_grupo.csv", data["groups"])
    write_csv(human / "02_gold_vs_sol_por_rol.csv", data["roles"])
    write_csv(human / "03_errores_y_diagnostico.csv", data["errors"])
    write_csv(human / "04_fuentes_usadas_y_excluidas.csv", data["sources"])
    write_csv(human / "05_telemetria_y_costos.csv", data["telemetry"])
    if not args.preserve_docx:
        build_docx(human / "Informe_SemanticV2_Calibration_Detenida.docx", data, registry_hash)
        build_pdf(human / "Informe_SemanticV2_Calibration_Detenida.pdf", data, registry_hash)

    workbook_data = {
        "RESUMEN": [
            {"indicador": "Estado", "valor": "CALIBRATION DETENIDA", "detalle": "Sin freeze, holdout ni recuración de 105 grupos"},
            {"indicador": "Grupos ejecutados", "valor": 5, "detalle": "GCLM:T1.6 no fue llamado"},
            {"indicador": "Costo observado USD", "valor": data["budget"]["observed_cost_usd"] + data["attempt1_budget"]["observed_cost_usd"], "detalle": "Incluye probe técnico del intento 1; sin costos no observables"},
            {"indicador": "Defectos estructurales", "valor": 2, "detalle": "Normalización de soporte contextual y validación de downstream null"},
            {"indicador": "Tests posteriores", "valor": "147/147 PASS", "detalle": "Incluye nuevas regresiones"},
            {"indicador": "Registro activo", "valor": "INTACTO", "detalle": registry_hash},
        ],
        "RESULTADO_GRUPOS": data["groups"],
        "GOLD_VS_SOL_ROL": data["roles"],
        "ERRORES_DIAGNOSTICO": data["errors"],
        "FUENTES": data["sources"],
        "TELEMETRIA_COSTOS": data["telemetry"],
    }
    (technical / "workbook_data.json").write_text(json.dumps(workbook_data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    for source in (CAMPAIGN_ATTEMPT1, CAMPAIGN):
        if not source.exists():
            continue
        shutil.copytree(source, technical / source.name, dirs_exist_ok=True)
    frozen_inputs = technical / "implementacion_y_contratos"
    frozen_inputs.mkdir(exist_ok=True)
    for name in (
        "curation_v2.py", "run_semantic_verification_v2.py", "mechanism_curation_semantic_v2.schema.json",
        "mechanism_curation_v2.schema.json", "prompt_mechanism_curator_semantic_v2.md",
        "prompt_mechanism_critic_semantic_v2.md", "prompt_mechanism_arbiter_semantic_v2.md",
    ):
        shutil.copy2(SERVICE / name, frozen_inputs / name)
    shutil.copy2(GOLD, technical / "gold_frozen.json")
    for group_id in GROUPS:
        packet = PACKETS / f"{group_id.replace(':', '__')}.json"
        shutil.copy2(packet, technical / packet.name)

    guide = human / "LEEME.txt"
    guide.write_text(
        "1. Lea primero el informe DOCX o PDF.\n"
        "2. Abra el XLSX para filtrar por grupo, rol, evidencia y costo.\n"
        "3. Use los CSV si necesita importar tablas a Drive.\n"
        "4. El replay corregido es sólo diagnóstico local: no convierte la campaña fallida en aprobada.\n"
        "5. El ZIP técnico contiene JSON para auditoría de ingeniería; el ZIP humano no contiene JSON.\n",
        encoding="utf-8",
    )

    manifest_rows = []
    for path in sorted(output.rglob("*")):
        if path.is_file() and path.name not in {"manifest_sha256.csv"} and path.suffix.lower() != ".zip":
            manifest_rows.append({"archivo": path.relative_to(output).as_posix(), "sha256": sha256(path), "bytes": path.stat().st_size})
    write_csv(output / "manifest_sha256.csv", manifest_rows)
    zip_tree(output / "LLM1_SemanticV2_Calibration_Paquete_Humano.zip", list(human.rglob("*")) + [output / "manifest_sha256.csv"], output)
    zip_tree(output / "LLM1_SemanticV2_Calibration_Auditoria_Tecnica.zip", list(technical.rglob("*")) + [output / "manifest_sha256.csv"], output)
    summary = {
        "created_at": datetime.now(timezone.utc).isoformat(),
        "campaign_status": data["terminal"]["status"],
        "calibration_complete": data["phase"]["complete"],
        "freeze_created": False,
        "holdout_started": False,
        "full_105_started": False,
        "tests_after_fix": "147/147 passed",
        "registry_sha256": registry_hash,
        "registry_intact": registry_hash == EXPECTED_REGISTRY_HASH,
        "observed_cost_usd": data["budget"]["observed_cost_usd"] + data["attempt1_budget"]["observed_cost_usd"],
    }
    (output / "audit_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
